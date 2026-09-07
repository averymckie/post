"""Fifth bounded adapter suite: P41-P50. Requires handoff_guards_v1.py, handoff_guards_v3.py and handoff_guards_v4.py from
this TXT. Every layout, pivot, join, decision evaluation, serialization, build, and comparison below is a library primitive
call; local code declares fixtures and policies and turns primitive results into Blocked outcomes. All runtime inputs are
authored fixtures. No network, model client, native Office application, or business action.
Run: python handoff_guards_v5.py --report report.json
"""
from __future__ import annotations
import argparse, base64, datetime, hashlib, importlib.metadata, io, json, logging, re, traceback, warnings
from pathlib import Path
from typing import Literal
import clingo
import docx
import networkx as nx
import numpy as np
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
import plotly.io as pio
import zen
from dateutil import parser as date_parser
from deepdiff import DeepDiff
from jinja2 import Environment, StrictUndefined
from pydantic import BaseModel, ConfigDict, Field, ValidationError, model_validator
from rapidfuzz import fuzz
import handoff_guards_v1 as g
import handoff_guards_v3 as g3
import handoff_guards_v4 as g4
warnings.simplefilter('ignore'); logging.disable(logging.WARNING)


def decode_typed_array(spec) -> np.ndarray:
    # plotly serializes numpy-backed arrays as {dtype, bdata, shape?}; numpy restores the buffer.
    if not isinstance(spec, dict): return np.asarray(spec, dtype=float)
    values = np.frombuffer(base64.b64decode(spec['bdata']), dtype=spec['dtype'])
    return values.reshape([int(x) for x in spec['shape'].split(',')]) if 'shape' in spec else values

# ---------------------------------------------------------------- P41
class StepNode(BaseModel):
    model_config = ConfigDict(strict=True, extra='forbid')
    id: str = Field(min_length=1)
    unit: int = Field(ge=0)


def layered_positions(nodes: list[dict], edges: list[tuple[str, str]]) -> dict:
    typed = [StepNode.model_validate(n) for n in nodes]
    graph = nx.DiGraph(); graph.add_nodes_from(n.id for n in typed); graph.add_edges_from(edges)
    if set(graph) != {n.id for n in typed}: raise g.Blocked('edge endpoint outside the step set')
    if not nx.is_directed_acyclic_graph(graph): raise g.Blocked('forced order is not acyclic')
    layers = {i: sorted(generation) for i, generation in enumerate(nx.topological_generations(graph))}
    pos = nx.multipartite_layout(graph, subset_key=layers, align='vertical')
    unit = {n.id: n.unit for n in typed}
    return {'coords': {n: [float(pos[n][0]), float(pos[n][1]), float(unit[n])] for n in sorted(graph)}, 'layers': len(layers)}


def process_3d_page(nodes: list[dict], edges: list[tuple[str, str]], *, include_plotlyjs) -> dict:
    layout = layered_positions(nodes, edges); coords = layout['coords']
    ids = sorted(coords)
    ex, ey, ez = [], [], []
    for a, b in edges:
        ex += [coords[a][0], coords[b][0], None]; ey += [coords[a][1], coords[b][1], None]; ez += [coords[a][2], coords[b][2], None]
    fig = go.Figure([go.Scatter3d(name='steps', x=[coords[i][0] for i in ids], y=[coords[i][1] for i in ids], z=[coords[i][2] for i in ids], text=ids, mode='markers+text'),
                     go.Scatter3d(name='forced', x=ex, y=ey, z=ez, mode='lines')])
    data = json.loads(pio.to_json(fig, validate=True))['data']
    diff = DeepDiff({'steps': [coords[i] for i in ids], 'edges': [ex, ey, ez]}, {'steps': [[x, y, z] for x, y, z in zip(data[0]['x'], data[0]['y'], data[0]['z'])], 'edges': [data[1]['x'], data[1]['y'], data[1]['z']]}, zip_ordered_iterables=True)
    if diff: raise g.Blocked('page payload differs from the layout: ' + diff.to_json())
    page = pio.to_html(fig, include_plotlyjs=include_plotlyjs, full_html=True, div_id='process3d', validate=True)
    external = g3.external_assets(page)
    if external: raise g.Blocked('external asset dependency: ' + external[0])
    return {'nodes': len(ids), 'edges': len(edges), 'layers': layout['layers'], 'bytes': len(page.encode())}

# ---------------------------------------------------------------- P42 and P47
def decision_table(inputs: list[tuple[str, str]], outputs: list[str], rules: list[dict]) -> dict:
    """A GoRules JDM graph as data: inputs are (field, kind) pairs, rules map field -> cell expression and output -> value."""
    inp = [{'id': 'i%d' % k, 'field': f, 'name': f, 'type': 'expression'} for k, (f, _) in enumerate(inputs)]
    out = [{'id': 'o%d' % k, 'field': f, 'name': f, 'type': 'expression'} for k, f in enumerate(outputs)]
    table = []
    for k, r in enumerate(rules):
        cells = {'_id': r['id']}
        for i in inp: cells[i['id']] = r['conditions'].get(i['field'], '')
        for o in out: cells[o['id']] = r['outputs'][o['field']]
        table.append(cells)
    return {'nodes': [{'id': 'in', 'type': 'inputNode', 'name': 'request', 'position': {'x': 0, 'y': 0}},
                      {'id': 'dt', 'type': 'decisionTableNode', 'name': 'policy', 'position': {'x': 0, 'y': 0}, 'content': {'hitPolicy': 'first', 'inputs': inp, 'outputs': out, 'rules': table}},
                      {'id': 'out', 'type': 'outputNode', 'name': 'response', 'position': {'x': 0, 'y': 0}}],
            'edges': [{'id': 'e1', 'type': 'edge', 'sourceId': 'in', 'targetId': 'dt'}, {'id': 'e2', 'type': 'edge', 'sourceId': 'dt', 'targetId': 'out'}]}


def load_decision(jdm: dict):
    return zen.ZenEngine().create_decision(json.dumps(jdm, allow_nan=False))


class MajorityInput(BaseModel):
    model_config = ConfigDict(strict=True, extra='forbid')
    present: int = Field(ge=0)

POLICY_PAGE = Environment(autoescape=True, undefined=StrictUndefined).from_string(
    '<table>{% for r in rows %}<tr><td>{{ r.id }}</td><td>{% for f, c in r.conditions %}{{ f }} {{ c }}; {% endfor %}</td><td>{% for f, v in r.outputs %}{{ f }} = {{ v }}; {% endfor %}</td><td>{{ r.note }}</td></tr>{% endfor %}</table>')


def policy_page(jdm: dict, samples: list[dict], input_model) -> str:
    decision = load_decision(jdm); table = jdm['nodes'][1]['content']
    rows = [{'id': r['_id'], 'conditions': [(i['field'], r[i['id']]) for i in table['inputs'] if r[i['id']]],
             'outputs': [(o['field'], r[o['id']]) for o in table['outputs']], 'note': r.get('_description', '')} for r in table['rules']]
    for s in samples:
        payload = input_model.model_validate(s['input']).model_dump()
        result = decision.evaluate(payload)['result']
        if result['rule'] != s['rule']: raise g.Blocked('page and engine disagree at sample %r: %r' % (s['input'], result))
    return POLICY_PAGE.render(rows=rows)

# ---------------------------------------------------------------- P43
def waiting_matrix(events: list[dict], *, departments: list[str], activities: list[str]) -> dict:
    df = pd.DataFrame(events)
    if df[['case_id', 'event_id']].duplicated().any(): raise g.Blocked('duplicate event identity')
    df['timestamp'] = pd.to_datetime(df['timestamp'], utc=True, errors='raise')
    if df.duplicated(['case_id', 'timestamp'], keep=False).any(): raise g.Blocked('timestamp ties need an ordering contract')
    if (df.groupby('case_id')['department'].nunique(dropna=False) > 1).any(): raise g.Blocked('department differs within a case')
    df = df.sort_values(['case_id', 'timestamp'], kind='stable')
    df['gap_hours'] = (df['timestamp'] - df.groupby('case_id')['timestamp'].shift(1)).dt.total_seconds() / 3600
    measured = df.dropna(subset=['gap_hours'])
    means = measured.pivot_table(index='department', columns='activity', values='gap_hours', aggfunc='mean').reindex(index=departments, columns=activities)
    counts = measured.pivot_table(index='department', columns='activity', values='gap_hours', aggfunc='count', fill_value=0).reindex(index=departments, columns=activities, fill_value=0)
    return {'means': means, 'counts': counts.astype(int), 'first_events_excluded': int(df['gap_hours'].isna().sum())}


def heatmap_page(matrix: dict, *, include_plotlyjs) -> dict:
    means, counts = matrix['means'], matrix['counts']
    fig = px.imshow(means, aspect='auto')
    fig.update_traces(customdata=counts.to_numpy(), hovertemplate='%{y} / %{x}: %{z} h over %{customdata} gaps<extra></extra>')
    trace = json.loads(pio.to_json(fig, validate=True))['data'][0]
    z = decode_typed_array(trace['z']); c = decode_typed_array(trace['customdata'])
    if not np.array_equal(z, means.to_numpy(dtype=float), equal_nan=True) or not np.array_equal(c, counts.to_numpy()): raise g.Blocked('heatmap payload differs from the matrix')
    if list(trace['x']) != list(means.columns) or list(trace['y']) != list(means.index): raise g.Blocked('heatmap axes differ from the declared axes')
    page = pio.to_html(fig, include_plotlyjs=include_plotlyjs, full_html=True, div_id='heatmap', validate=True)
    external = g3.external_assets(page)
    if external: raise g.Blocked('external asset dependency: ' + external[0])
    return {'cells': int(z.size), 'missing_cells': int(np.isnan(z).sum()), 'bytes': len(page.encode())}

# ---------------------------------------------------------------- P44
ACTOR_LEXICON = {'AGENCY': ['agency', 'each agency', 'the agency'], 'TSC': ['tsc', 'the tsc']}


def actor_matrix(facts: list[dict], *, key: str, actions: list[str]) -> pd.DataFrame:
    df = pd.DataFrame(facts); df = df[df['modality'] == 'obligatory']
    return df.pivot_table(index=key, columns='lemma', values='event', aggfunc='count', fill_value=0).reindex(columns=actions, fill_value=0)


def concept_actor_matrix(facts: list[dict], *, actions: list[str]) -> dict:
    rows = []
    for f in facts:
        if f['modality'] != 'obligatory': continue
        concept = g.resolve_alias(f['agent_lemma'], ACTOR_LEXICON, scope=f['scope'], allowed_scope='obligation', polarity=f['polarity'], expected_polarity='positive')
        rows.append(dict(f, actor=concept))
    matrix = actor_matrix(rows, key='actor', actions=actions)
    if int(matrix.to_numpy().sum()) != len(rows): raise g.Blocked('matrix does not conserve the obligatory facts')
    return {'actors': list(matrix.index), 'matrix': {a: [int(v) for v in matrix.loc[a]] for a in matrix.index}, 'facts': len(rows)}

# ---------------------------------------------------------------- P45
MEETING_ID = re.compile(r'tsc-(\d{4}-\d{2}-\d{2})')


def meeting_date(meeting: str) -> datetime.date:
    match = MEETING_ID.fullmatch(meeting)
    if not match: raise g.Blocked('meeting identifier does not carry a full calendar date: ' + meeting)
    return datetime.date.fromisoformat(match.group(1))


def timeline_page(tags: list[dict], *, include_plotlyjs) -> dict:
    df = pd.DataFrame(tags)
    if set(df['state']) - {'accepted', 'tie'}: raise g.Blocked('unknown tag state')
    accepted = df[df['state'] == 'accepted']
    points = accepted.groupby(['meeting', 'step'], dropna=False).size().reset_index(name='lines')
    points['date'] = [meeting_date(m).isoformat() for m in points['meeting']]
    fig = go.Figure(go.Scatter(x=points['date'].tolist(), y=points['step'].tolist(), mode='markers', marker={'size': points['lines'].tolist()},
                               customdata=[[int(n)] for n in points['lines']], hovertemplate='%{y} on %{x}: %{customdata[0]} lines<extra></extra>'))
    trace = json.loads(pio.to_json(fig, validate=True))['data'][0]
    diff = DeepDiff({'x': points['date'].tolist(), 'y': points['step'].tolist(), 'customdata': [[int(n)] for n in points['lines']]}, {k: trace[k] for k in ['x', 'y', 'customdata']}, zip_ordered_iterables=True)
    if diff: raise g.Blocked('timeline payload differs from the counts: ' + diff.to_json())
    page = pio.to_html(fig, include_plotlyjs=include_plotlyjs, full_html=True, div_id='timeline', validate=True)
    external = g3.external_assets(page)
    if external: raise g.Blocked('external asset dependency: ' + external[0])
    return {'points': len(points), 'lines_total': int(points['lines'].sum()), 'tied_excluded': int((df['state'] == 'tie').sum()), 'bytes': len(page.encode())}

# ---------------------------------------------------------------- P46
def edge_waits(events: list[dict]) -> list[dict]:
    df = pd.DataFrame(events)
    if df[['case_id', 'event_id']].duplicated().any(): raise g.Blocked('duplicate event identity')
    df['timestamp'] = pd.to_datetime(df['timestamp'], utc=True, errors='raise')
    if df.duplicated(['case_id', 'timestamp'], keep=False).any(): raise g.Blocked('timestamp ties need an ordering contract')
    df = df.sort_values(['case_id', 'timestamp'], kind='stable')
    df['previous'] = df.groupby('case_id')['activity'].shift(1)
    df['gap_hours'] = (df['timestamp'] - df.groupby('case_id')['timestamp'].shift(1)).dt.total_seconds() / 3600
    edges = df.dropna(subset=['previous']).groupby(['previous', 'activity'])['gap_hours'].agg(cases='count', mean_hours='mean').reset_index()
    edges = edges.sort_values(['mean_hours', 'cases', 'previous', 'activity'], ascending=[False, False, True, True], kind='stable')
    return [dict(r, measure='edge_wait') for r in g3.records(edges)]

# ---------------------------------------------------------------- P47
class OpenCase(BaseModel):
    model_config = ConfigDict(strict=True, extra='forbid')
    case_id: str = Field(min_length=1)
    department: str = Field(min_length=1)
    elapsed_days: float = Field(ge=0)


def at_risk_decisions(department_means: dict, open_cases: list[dict]) -> dict:
    rules = [{'id': 'r_' + d, 'conditions': {'department': json.dumps(d), 'elapsed_days': '> %r' % float(m) if m is not None else '> nan'}, 'outputs': {'at_risk': 'true', 'rule': json.dumps('r_' + d)}} for d, m in department_means.items()]
    rules.append({'id': 'r_default', 'conditions': {}, 'outputs': {'at_risk': 'false', 'rule': '"r_default"'}})
    for d, m in department_means.items():
        if m is None or not np.isfinite(m): raise g.Blocked('department without a finite closed-case mean: ' + d)
    jdm = decision_table([('department', 'string'), ('elapsed_days', 'number')], ['at_risk', 'rule'], rules)
    decision = load_decision(jdm)
    for d, m in department_means.items():  # boundary samples prove the comparison the table encodes
        if decision.evaluate({'department': d, 'elapsed_days': float(m)})['result']['at_risk'] is not False: raise g.Blocked('boundary is not exclusive for ' + d)
        if decision.evaluate({'department': d, 'elapsed_days': float(m) + 1e-9})['result']['at_risk'] is not True: raise g.Blocked('threshold does not fire for ' + d)
    out = []
    for c in open_cases:
        typed = OpenCase.model_validate(c)
        if typed.department not in department_means: raise g.Blocked('department without a threshold: ' + typed.department)
        out.append({'case_id': typed.case_id, **decision.evaluate(typed.model_dump(exclude={'case_id'}))['result']})
    return {'decisions': out, 'basis': 'threshold derived from closed-case means of the same log; not a policy', 'jdm_sha256': hashlib.sha256(g.canonical_local(jdm)).hexdigest()}

# ---------------------------------------------------------------- P48
class ChecklistItem(BaseModel):
    model_config = ConfigDict(strict=True, extra='forbid')
    item_id: str = Field(min_length=1)
    fact_id: str = Field(min_length=1)
    sentence: str = Field(min_length=1)
    quote: str = Field(min_length=1)

    @model_validator(mode='after')
    def quote_inside(self):
        if self.quote not in self.sentence: raise ValueError('quote is not inside its sentence')
        return self

BOX = '☐'


def checklist_document(items: list[dict]) -> bytes:
    typed = [ChecklistItem.model_validate(i) for i in items]
    if pd.Index([i.item_id for i in typed]).has_duplicates: raise g.Blocked('duplicate item identity')
    d = docx.Document()
    for i in typed:
        p = d.add_paragraph(); p.add_run(BOX + ' '); hidden = p.add_run(i.item_id + '|' + i.fact_id + '|' + i.quote); hidden.font.hidden = True; p.add_run(i.sentence)
    buf = io.BytesIO(); d.save(buf); return buf.getvalue()


def read_checklist(data: bytes) -> list[dict]:
    out = []
    for p in docx.Document(io.BytesIO(data)).paragraphs:
        runs = p.runs
        if len(runs) != 3 or runs[0].text != BOX + ' ' or runs[1].font.hidden is not True: raise g.Blocked('checklist paragraph lost its identity run')
        item_id, fact_id, quote = runs[1].text.split('|')
        out.append(ChecklistItem(item_id=item_id, fact_id=fact_id, sentence=runs[2].text, quote=quote).model_dump())
    return out

# ---------------------------------------------------------------- P49
CROSS_PROGRAM = """shared_lemma(L) :- action(a, _, L, _), action(b, _, L, _).
shared_concept(C) :- action(a, _, _, C), action(b, _, _, C).
only_lemma(D, L) :- action(D, _, L, _), other(D, E), not action(E, _, L, _).
only_concept(D, C) :- action(D, _, _, C), other(D, E), not action(E, _, _, C).
other(a, b). other(b, a).
#show shared_lemma/1. #show shared_concept/1. #show only_lemma/2. #show only_concept/2."""


def cross_document(actions_a: list[dict], actions_b: list[dict]) -> dict:
    ctl = clingo.Control(['0', '--warn=no-atom-undefined'])
    with ctl.backend() as be:
        for doc, rows in (('a', actions_a), ('b', actions_b)):
            for r in rows:
                be.add_rule([be.add_atom(clingo.Function('action', [clingo.Function(doc), clingo.String(r['id']), clingo.String(r['lemma']), clingo.String(r['concept'])]))])
    ctl.add('base', [], CROSS_PROGRAM); ctl.ground([('base', [])])
    shown = []; ctl.solve(on_model=lambda m: shown.extend(m.symbols(shown=True)))
    out = {'shared_lemma': [], 'shared_concept': [], 'only_lemma': [], 'only_concept': []}
    for s in shown:
        out[s.name].append(s.arguments[0].string if len(s.arguments) == 1 else [s.arguments[0].name, s.arguments[1].string])
    return {k: sorted(v) for k, v in out.items()}

# ---------------------------------------------------------------- P50
def site_steps_roundtrip(steps: list[dict]) -> dict:
    rows = [g4.StepRow.model_validate(s).model_dump() for s in steps]
    if pd.Index([r['id'] for r in rows]).has_duplicates: raise g.Blocked('duplicate step identity')
    page = g4.build_site(g4.table_markdown(rows, g4.STEP_COLUMNS), highlightjs=False)
    external = g3.external_assets(page)
    if external: raise g.Blocked('external asset dependency: ' + external[0])
    back = [dict(zip(g4.STEP_COLUMNS, cells)) for cells in g4.rendered_rows(page, '//*[@role="main"]//table[1]')]
    forward, backward = hashlib.sha256(g.canonical_local(rows)).hexdigest(), hashlib.sha256(g.canonical_local(back)).hexdigest()
    diff = DeepDiff(rows, back, zip_ordered_iterables=True)
    if diff or forward != backward: raise g.Blocked('site readback differs: ' + (diff.to_json() if diff else 'digest'))
    return {'rows': len(back), 'digest': forward, 'profile': 'canonical_local JSON of typed rows', 'all_tables_rows': len(g4.rendered_rows(page, '//table'))}

C = []
def case(i):
    def reg(fn): C.append((i, fn)); return fn
    return reg


@case(41)
def deterministic_layers_and_payload():
    edges = [('request', 'make'), ('receipt', 'except'), ('determination', 'make')]
    first = nx.DiGraph(); first.add_nodes_from(['request', 'receipt', 'determination', 'make', 'except']); first.add_edges_from(edges)
    second = nx.DiGraph(); second.add_nodes_from(['determination', 'make', 'except', 'receipt', 'request']); second.add_edges_from(edges)
    g.equal([list(x) for x in nx.topological_generations(first)], [['request', 'receipt', 'determination'], ['except', 'make']])
    g.equal([list(x) for x in nx.topological_generations(second)], [['determination', 'receipt', 'request'], ['except', 'make']])  # insertion order leaks into the layer order
    nodes = [{'id': 'request', 'unit': 1}, {'id': 'receipt', 'unit': 1}, {'id': 'determination', 'unit': 2}, {'id': 'make', 'unit': 1}, {'id': 'except', 'unit': 3}]
    g.equal(bool(DeepDiff(layered_positions(nodes, edges), layered_positions(list(reversed(nodes)), edges))), False)
    out = process_3d_page(nodes, edges, include_plotlyjs=True)
    g.equal((out['nodes'], out['edges'], out['layers']), (5, 3, 2))
    encoded = json.loads(pio.to_json(go.Figure(go.Scatter3d(x=np.array([0.0, 1.0, np.nan]), y=[0, 0, None], z=[0, 0, None], mode='lines'))))['data'][0]
    g.equal(bool(np.isnan(decode_typed_array(encoded['x'])[2])), True); g.equal(encoded['y'], [0, 0, None])  # the numpy path hides the separator as NaN
    g.rejects(g.Blocked, lambda: process_3d_page(nodes, edges, include_plotlyjs='cdn'))
    g.rejects(g.Blocked, lambda: layered_positions(nodes, edges + [('make', 'request')]))
    g.rejects(g.Blocked, lambda: layered_positions(nodes, edges + [('make', 'zz')]))
    return {'insertion_order_in_generations_reproduced': True, 'positions_deterministic_under_reordering': True, 'edge_separators_kept_as_null': True,
            'unit_index_carried_as_z': True, 'cdn_page': 'blocked', 'original_chain_rerun': False}


@case(42)
def policy_page_from_structure_with_engine_agreement():
    rules = [{'id': 'r1', 'conditions': {'present': '>= 10'}, 'outputs': {'majority_reachable': 'true', 'rule': '"r1"'}},
             {'id': 'r2', 'conditions': {}, 'outputs': {'majority_reachable': 'false', 'rule': '"r2"'}}]
    jdm = decision_table([('present', 'number')], ['majority_reachable', 'rule'], rules)
    jdm['nodes'][1]['content']['rules'][0]['_description'] = '<script>alert(1)</script> if present >= 10'
    decision = load_decision(jdm)
    g.equal([decision.evaluate({'present': v})['result']['rule'] for v in (9, 10, '10', None)], ['r2', 'r1', 'r2', 'r2'])  # a string or missing value takes the fallback rule silently
    g.equal(decision.evaluate({})['result']['rule'], 'r2')
    samples = [{'input': {'present': 9}, 'rule': 'r2'}, {'input': {'present': 10}, 'rule': 'r1'}]
    page = policy_page(jdm, samples, MajorityInput)
    g.equal('&lt;script&gt;' in page and '<script>' not in page, True); g.equal('present &gt;= 10' in page, True)
    unsafe = Environment(autoescape=False).from_string('{{ note }}').render(note=jdm['nodes'][1]['content']['rules'][0]['_description'])
    g.equal('<script>' in unsafe, True)
    g.rejects(g.Blocked, lambda: policy_page(jdm, [{'input': {'present': 9}, 'rule': 'r1'}], MajorityInput))
    g.rejects(ValidationError, lambda: policy_page(jdm, [{'input': {'present': '10'}, 'rule': 'r1'}], MajorityInput))
    g.rejects(ValidationError, lambda: policy_page(jdm, [{'input': {}, 'rule': 'r2'}], MajorityInput))
    return {'fallback_on_missing_or_string_input_reproduced': True, 'autoescape_off_injection_reproduced': True, 'page_rendered_from_table_structure': True,
            'engine_agreement_at_samples': True, 'typed_inputs_before_evaluate': True, 'original_chain_rerun': False}


@case(43)
def heatmap_cells_with_partitioned_lags():
    t = lambda h: '2011-10-01T%02d:00:00Z' % h
    events = [{'case_id': 'c1', 'event_id': 'e1', 'department': 'General', 'activity': 'T02', 'timestamp': t(0)},
              {'case_id': 'c1', 'event_id': 'e2', 'department': 'General', 'activity': 'T03', 'timestamp': t(4)},
              {'case_id': 'c2', 'event_id': 'e3', 'department': 'Experts', 'activity': 'T02', 'timestamp': t(1)},
              {'case_id': 'c2', 'event_id': 'e4', 'department': 'Experts', 'activity': 'T03', 'timestamp': t(11)}]
    unpartitioned = pd.DataFrame(events); unpartitioned['timestamp'] = pd.to_datetime(unpartitioned['timestamp'], utc=True)
    unpartitioned = unpartitioned.sort_values('timestamp'); crossed = unpartitioned['timestamp'].diff().dt.total_seconds() / 3600
    g.equal(crossed.tolist()[1:], [1.0, 3.0, 7.0])  # a lag without case partition measures across cases
    matrix = waiting_matrix(events, departments=['Experts', 'General', 'Customer contact'], activities=['T02', 'T03'])
    g.equal(matrix['means'].loc['General', 'T03'], 4.0); g.equal(matrix['means'].loc['Experts', 'T03'], 10.0)
    g.equal(bool(np.isnan(matrix['means'].loc['Customer contact', 'T03'])), True); g.equal(int(matrix['counts'].loc['Customer contact', 'T03']), 0)
    g.equal(matrix['first_events_excluded'], 2)
    out = heatmap_page(matrix, include_plotlyjs=True)
    g.equal((out['cells'], out['missing_cells']), (6, 4))
    fig = px.imshow(matrix['means']); back = pio.from_json(pio.to_json(fig))
    g.equal(sorted(back.data[0].z.keys()) if isinstance(back.data[0].z, dict) else 'array', ['bdata', 'dtype', 'shape'])  # from_json returns the typed buffer undecoded
    g.rejects(g.Blocked, lambda: heatmap_page(matrix, include_plotlyjs='cdn'))
    g.rejects(g.Blocked, lambda: waiting_matrix(events + [dict(events[1], event_id='e9', department='Experts')], departments=['Experts', 'General'], activities=['T02', 'T03']))
    return {'cross_case_lag_reproduced': True, 'missing_cells_explicit_with_zero_counts': True, 'payload_decoded_and_compared': True,
            'from_json_leaves_buffer_encoded_reproduced': True, 'cdn_page': 'blocked', 'original_chain_rerun': False}


@case(44)
def actor_axis_by_identity_not_quote():
    facts = [{'event': 'e7', 'agent_arg': 'x5', 'agent_lemma': 'agency', 'agent_quote': 'agency shall make', 'lemma': 'make', 'modality': 'obligatory', 'scope': 'obligation', 'polarity': 'positive'},
             {'event': 'e9', 'agent_arg': 'x8', 'agent_lemma': 'agency', 'agent_quote': 'Each agency shall notify', 'lemma': 'notify', 'modality': 'obligatory', 'scope': 'obligation', 'polarity': 'positive'},
             {'event': 'e12', 'agent_arg': 'x11', 'agent_lemma': 'agency', 'agent_quote': 'the agency shall promulgate', 'lemma': 'promulgate', 'modality': 'obligatory', 'scope': 'obligation', 'polarity': 'positive'},
             {'event': 'e3', 'agent_arg': 'x2', 'agent_lemma': 'TSC', 'agent_quote': 'TSC must have', 'lemma': 'have', 'modality': 'obligatory', 'scope': 'obligation', 'polarity': 'positive'},
             {'event': 'e4', 'agent_arg': 'x3', 'agent_lemma': 'TSC', 'agent_quote': 'TSC may meet', 'lemma': 'meet', 'modality': 'permitted', 'scope': 'obligation', 'polarity': 'positive'}]
    actions = ['make', 'notify', 'promulgate', 'have']
    g.equal(len(actor_matrix(facts, key='agent_quote', actions=actions)), 4)  # one row per quote: three rows for one agency
    g.equal(len(actor_matrix(facts, key='agent_lemma', actions=actions)), 2)
    out = concept_actor_matrix(facts, actions=actions)
    g.equal(out['matrix'], {'AGENCY': [1, 1, 1, 0], 'TSC': [0, 0, 0, 1]}); g.equal(out['facts'], 4)
    g.rejects(g.Blocked, lambda: concept_actor_matrix(facts + [dict(facts[0], event='e99', agent_lemma='board')], actions=actions))
    g.rejects(g.Blocked, lambda: concept_actor_matrix([dict(facts[0], polarity='negative')], actions=actions))
    return {'quote_keyed_axis_fragmentation_reproduced': True, 'actor_identity_through_approved_lexicon': True, 'permitted_facts_excluded': True,
            'conservation_checked': True, 'original_chain_rerun': False}


@case(45)
def timeline_dates_from_declared_pattern():
    fuzzy = date_parser.parse('tsc-2024-01', fuzzy=True)
    g.equal((fuzzy.year, fuzzy.month, fuzzy.day == datetime.date.today().day), (2024, 1, True))  # the missing day comes from today
    g.equal(date_parser.parse('tsc-2024-01-17', fuzzy=True).date(), datetime.date(2024, 1, 17))
    g.equal(meeting_date('tsc-2024-01-17'), datetime.date(2024, 1, 17))
    g.rejects(g.Blocked, lambda: meeting_date('tsc-2024-01'))
    g.rejects(ValueError, lambda: meeting_date('tsc-2024-13-40'))
    tags = [{'meeting': 'tsc-2024-01-17', 'step': 'share', 'state': 'accepted'}, {'meeting': 'tsc-2024-01-17', 'step': 'share', 'state': 'accepted'},
            {'meeting': 'tsc-2024-01-17', 'step': 'meeting', 'state': 'tie'}, {'meeting': 'tsc-2023-11-08', 'step': 'ensure', 'state': 'accepted'}]
    out = timeline_page(tags, include_plotlyjs=True)
    g.equal((out['points'], out['lines_total'], out['tied_excluded']), (2, 3, 1))
    g.rejects(g.Blocked, lambda: timeline_page(tags + [{'meeting': 'tsc-2025', 'step': 'x', 'state': 'accepted'}], include_plotlyjs=True))
    g.rejects(g.Blocked, lambda: timeline_page(tags, include_plotlyjs='cdn'))
    return {'fuzzy_parse_invents_day_reproduced': True, 'declared_pattern_and_fromisoformat': True, 'counts_in_customdata_not_only_marker_size': True,
            'ties_excluded_and_reported': True, 'original_chain_rerun': False}


@case(46)
def edge_waits_not_target_means():
    edges = pd.DataFrame({'source': ['T05', 'T02', 'T04'], 'target': ['T13', 'T03', 'T03'], 'cases': [2, 43, 6]})
    activity_wait = pd.DataFrame({'activity': ['T13', 'T03'], 'mean_hours': [205.36, 129.66]})
    joined = edges.merge(activity_wait, left_on='target', right_on='activity', validate='many_to_one')
    g.equal(joined.loc[joined['target'] == 'T03', 'mean_hours'].tolist(), [129.66, 129.66])  # every edge into T03 inherits the same activity mean
    t = lambda h: '2011-10-01T%02d:00:00Z' % h
    events = [{'case_id': 'c1', 'event_id': 'e1', 'activity': 'T02', 'timestamp': t(0)}, {'case_id': 'c1', 'event_id': 'e2', 'activity': 'T03', 'timestamp': '2011-10-02T00:00:00Z'},
              {'case_id': 'c2', 'event_id': 'e3', 'activity': 'T04', 'timestamp': t(0)}, {'case_id': 'c2', 'event_id': 'e4', 'activity': 'T03', 'timestamp': t(12)}]
    out = edge_waits(events)
    g.equal([(r['previous'], r['activity'], r['cases'], r['mean_hours'], r['measure']) for r in out], [('T02', 'T03', 1, 24.0, 'edge_wait'), ('T04', 'T03', 1, 12.0, 'edge_wait')])
    g.rejects(g.Blocked, lambda: edge_waits(events + [dict(events[3], event_id='e5', activity='T09')]))
    return {'target_mean_duplicated_across_edges_reproduced': True, 'per_edge_wait_from_partitioned_shift': True, 'ranking_with_declared_tiebreak': True,
            'measure_labeled': True, 'original_chain_rerun': False}


@case(47)
def derived_thresholds_need_finite_values_and_typed_inputs():
    rendered = Environment(undefined=StrictUndefined).from_string('{"threshold": {{ mean }}}').render(mean=float('nan'))
    g.rejects(Exception, lambda: json.loads(rendered))  # text templating writes nan into JSON
    g.equal(Environment().from_string('{{ v | tojson }}').render(v=float('nan')), 'NaN')
    g.rejects(g.Blocked, lambda: g.strict_json_loads('{"threshold": NaN}'))
    nan_rules = [{'id': 'r_General', 'conditions': {'department': '"General"', 'elapsed_days': '> nan'}, 'outputs': {'at_risk': 'true', 'rule': '"r_General"'}},
                 {'id': 'r_default', 'conditions': {}, 'outputs': {'at_risk': 'false', 'rule': '"r_default"'}}]
    silent = load_decision(decision_table([('department', 'string'), ('elapsed_days', 'number')], ['at_risk', 'rule'], nan_rules))
    g.equal(silent.evaluate({'department': 'General', 'elapsed_days': 60.0})['result']['rule'], 'r_default')  # a nan threshold never fires and raises nothing
    g.equal(silent.evaluate({'department': 'General', 'elapsed_days': '60'})['result']['rule'], 'r_default')
    means = {'General': 5.42, 'Experts': 11.32}
    out = at_risk_decisions(means, [{'case_id': 'case-10011', 'department': 'General', 'elapsed_days': 60.0}, {'case_id': 'case-10012', 'department': 'General', 'elapsed_days': 5.42}])
    g.equal([(d['case_id'], d['at_risk'], d['rule']) for d in out['decisions']], [('case-10011', True, 'r_General'), ('case-10012', False, 'r_default')])
    g.rejects(g.Blocked, lambda: at_risk_decisions({'General': 5.42, 'Customer contact': None}, []))
    g.rejects(g.Blocked, lambda: at_risk_decisions({'General': float('nan')}, []))
    g.rejects(ValidationError, lambda: at_risk_decisions(means, [{'case_id': 'x', 'department': 'General', 'elapsed_days': '60'}]))
    g.rejects(g.Blocked, lambda: at_risk_decisions(means, [{'case_id': 'x', 'department': 'Post', 'elapsed_days': 1.0}]))
    return {'nan_through_text_template_reproduced': True, 'tojson_emits_bare_nan_reproduced': True, 'nan_threshold_silently_never_fires_reproduced': True,
            'boundary_samples_prove_strict_comparison': True, 'basis_labeled_not_policy': True, 'original_chain_rerun': False}


@case(48)
def checklist_items_carry_identity_and_no_completion():
    items = [{'item_id': 'k1', 'fact_id': 'tsc:e3', 'sentence': 'The TSC must have a Chair.', 'quote': 'must have'},
             {'item_id': 'k2', 'fact_id': 'tsc:e4', 'sentence': 'The TSC must have a Chair.', 'quote': 'must have'}]
    data = checklist_document(items)
    g.equal(read_checklist(data), items)
    d = docx.Document(io.BytesIO(data)); g.equal([len(p.runs) for p in d.paragraphs], [3, 3])
    d.paragraphs[0].text = d.paragraphs[0].text  # assigning text collapses the runs
    g.equal(len(d.paragraphs[0].runs), 1); buf = io.BytesIO(); d.save(buf)
    g.rejects(g.Blocked, lambda: read_checklist(buf.getvalue()))
    g.rejects(g.Blocked, lambda: g.observed_action('checked_glyph', {'case_id': 'c', 'action_id': 'k1', 'receipt_id': 'r', 'timestamp': 't'}))  # a glyph is not an occurrence receipt
    g.equal(g.observed_action('observed_completion', {'case_id': 'c', 'action_id': 'k1', 'receipt_id': 'r', 'timestamp': '2026-01-01T00:00:00Z'}), 'observed')
    g.rejects(g.Blocked, lambda: checklist_document(items + [items[0]]))
    g.rejects(ValidationError, lambda: checklist_document([dict(items[0], quote='shall')]))
    return {'identity_in_hidden_run_round_trips': True, 'text_assignment_drops_runs_reproduced': True, 'glyph_state_is_not_completion': True,
            'duplicate_identity_blocked': True, 'original_chain_rerun': False}


@case(49)
def cross_document_join_on_concepts():
    charter = [{'id': 'c1', 'lemma': 'approve', 'concept': 'APPROVE_MEMBERSHIP'}, {'id': 'c2', 'lemma': 'have', 'concept': 'HAVE_CHAIR'}, {'id': 'c3', 'lemma': 'nominate', 'concept': 'NOMINATE_MEMBER'}]
    governance = [{'id': 'g1', 'lemma': 'approve', 'concept': 'APPROVE_PULL_REQUEST'}, {'id': 'g2', 'lemma': 'nominat', 'concept': 'NOMINATE_MEMBER'}]
    out = cross_document(charter, governance)
    g.equal(out['shared_lemma'], ['approve']); g.equal(out['shared_concept'], ['NOMINATE_MEMBER'])  # the lemma join shares the wrong action and misses the shared one
    g.equal(out['only_lemma'], [['a', 'have'], ['a', 'nominate'], ['b', 'nominat']])
    g.equal(out['only_concept'], [['a', 'APPROVE_MEMBERSHIP'], ['a', 'HAVE_CHAIR'], ['b', 'APPROVE_PULL_REQUEST']])
    g.equal(round(fuzz.ratio('nominat', 'nominate')) > 90, True)  # a lemmatizer artifact is a near candidate, not an identity
    return {'lemma_join_conflation_reproduced': True, 'lemma_join_miss_on_artifact_reproduced': True, 'concept_join_from_approved_ids': True,
            'grounded_through_backend_symbols': True, 'original_chain_rerun': False}


@case(50)
def site_readback_selects_the_content_table():
    steps = [{'id': 'u0131:s00#e16', 'lemma': 'authorize', 'quote': 'authorize | delegate'}, {'id': 'u0089:s02#e53', 'lemma': 'make', 'quote': '<b>shall make</b>'}]
    out = site_steps_roundtrip(steps)
    g.equal(out['rows'], 2); g.equal(out['all_tables_rows'] > out['rows'], True)  # //table also matches theme tables
    g.equal(out['digest'], hashlib.sha256(g.canonical_local(steps)).hexdigest())
    g.rejects(g.Blocked, lambda: site_steps_roundtrip(steps + [steps[0]]))
    g.rejects(ValidationError, lambda: site_steps_roundtrip([dict(steps[0], lemma='')]))
    return {'theme_tables_in_generic_xpath_reproduced': True, 'content_table_by_role': True, 'digest_over_declared_profile': True,
            'escaped_cells_round_trip': True, 'original_chain_rerun': False}


def run() -> dict:
    g.ASSERTIONS = 0; rows = []
    for i, fn in C:
        before = g.ASSERTIONS
        try: r = fn(); state = 'passed'; error = None
        except Exception: r = {}; state = 'failed'; error = traceback.format_exc()
        rows.append({'proof': i, 'test': fn.__name__, 'status': state, 'assertions': g.ASSERTIONS - before, 'result': r, 'error': error})
    packages = ['pydantic', 'pandas', 'numpy', 'deepdiff', 'plotly', 'networkx', 'zen-engine', 'Jinja2', 'python-dateutil', 'python-docx', 'clingo', 'rapidfuzz', 'spacy', 'mkdocs', 'Markdown', 'lxml']
    here = Path(__file__).resolve().parent
    return {'scope': 'Bounded authored handoff fixtures and reproduced library defaults; the original P41-P50 chains and their artifacts were not rerun.',
            'test_cases': len(rows), 'passed': sum(r['status'] == 'passed' for r in rows), 'failed': sum(r['status'] == 'failed' for r in rows),
            'assertions': g.ASSERTIONS, 'full_chains_executed': 0, 'network_calls': 0, 'business_actions': 0,
            'source_sha256': hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
            'dependency_sha256': {name: hashlib.sha256((here / name).read_bytes()).hexdigest() for name in ['handoff_guards_v1.py', 'handoff_guards_v3.py', 'handoff_guards_v4.py']},
            'library_versions': {n: importlib.metadata.version(n) for n in packages}, 'results': rows}


if __name__ == '__main__':
    p = argparse.ArgumentParser(); p.add_argument('--report', type=Path, required=True); args = p.parse_args()
    result = run(); args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(json.dumps(result, indent=2, sort_keys=True, allow_nan=False) + '\n')
    print(json.dumps({k: result[k] for k in ['test_cases', 'passed', 'failed', 'assertions', 'full_chains_executed']}))
    if result['failed']:
        print('\n'.join(r['error'] for r in result['results'] if r['error'])); raise SystemExit(1)
