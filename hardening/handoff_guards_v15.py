"""Fifteenth bounded adapter suite: P141-P150, project schedule, cost and cash handoffs. Requires handoff_guards_v1.py
from this TXT. Every interval, working-day count, graph check, join, rollup, bucket, exact-decimal aggregation and test
report reading below is a library primitive call; local code declares fixtures and policies and turns a primitive result
into a Blocked outcome. All runtime inputs are authored fixtures. No model client and no network.
Run: python handoff_guards_v15.py --report report.json
"""
from __future__ import annotations
import argparse, datetime, hashlib, importlib.metadata, io, itertools, json, logging, subprocess, sys, tempfile, traceback, warnings
from decimal import Decimal, InvalidOperation
from fractions import Fraction
from pathlib import Path
import duckdb
import matplotlib
matplotlib.use('Agg')
import matplotlib.dates as mdates
import matplotlib.pyplot as plt
import networkx as nx
import numpy as np
import openpyxl
import pandas as pd
import xlsxwriter
from defusedxml import ElementTree as DET
import handoff_guards_v1 as g
warnings.simplefilter('ignore'); logging.disable(logging.WARNING)

CALENDAR = np.busdaycalendar(weekmask='Mon Tue Wed Thu Fri', holidays=['2026-09-07'])

# ---------------------------------------------------------------- P141
def bar_spans(intervals: list[dict], *, width_from: str) -> list[tuple[float, float]]:
    """width_from='width' passes (start, finish-start); width_from='finish' passes (start, finish), the common error."""
    out = []
    for i in intervals:
        start = mdates.date2num(datetime.date.fromisoformat(i['start']))
        finish = mdates.date2num(datetime.date.fromisoformat(i['finish']))
        if finish < start: raise g.Blocked('finish before start for %s' % i['id'])
        out.append((start, finish - start) if width_from == 'width' else (start, finish))
    return out


def drawn_widths(spans: list[tuple[float, float]]) -> list[float]:
    fig, ax = plt.subplots()
    try:
        return [float(ax.broken_barh([s], (n, 0.8)).get_paths()[0].get_extents().width) for n, s in enumerate(spans)]
    finally:
        plt.close(fig)


def schedule_table(rows: list[dict], columns: list[str], *, declare_columns: bool) -> bytes:
    buf = io.BytesIO(); wb = xlsxwriter.Workbook(buf, {'in_memory': True}); ws = wb.add_worksheet('schedule')
    ws.write_row(0, 0, columns)
    for n, r in enumerate(rows, start=1): ws.write_row(n, 0, [r[c] for c in columns])
    options = {'columns': [{'header': c} for c in columns]} if declare_columns else None
    code = ws.add_table(0, 0, len(rows), len(columns) - 1, options)
    wb.close()
    if code != 0: raise g.Blocked('add_table refused the range and returned %d' % code)
    return buf.getvalue()


def workbook_rows(data: bytes) -> list[list]:
    ws = openpyxl.load_workbook(io.BytesIO(data))['schedule']
    return [[c.value for c in row] for row in ws.iter_rows()]

# ---------------------------------------------------------------- P142
def working_days(start: str, finish: str, *, inclusive: bool) -> int:
    end = np.datetime64(finish) + (1 if inclusive else 0)
    return int(np.busday_count(np.datetime64(start), end, busdaycal=CALENDAR))


def offset_day(day: str, days: int, *, roll: str) -> str:
    return str(np.busday_offset(np.datetime64(day), days, roll=roll, busdaycal=CALENDAR))


def build_plan(tasks: list[dict]) -> nx.DiGraph:
    declared = {t['id'] for t in tasks}
    graph = nx.DiGraph()
    graph.add_nodes_from((t['id'], {'days': t['days']}) for t in tasks)
    for t in tasks:
        for p in t['after']:
            if p not in declared: raise g.Blocked('undeclared predecessor %r for %s' % (p, t['id']))
            graph.add_edge(p, t['id'])
    if not nx.is_directed_acyclic_graph(graph): raise g.Blocked('the plan has a cycle')
    return graph


def earliest_starts(graph: nx.DiGraph) -> dict:
    out = {}
    for node in nx.lexicographical_topological_sort(graph):
        out[node] = max((out[p] + graph.nodes[p]['days'] for p in graph.predecessors(node)), default=0)
    return out

# ---------------------------------------------------------------- P143
def wbs_graph(tasks: list[dict], packages: list[dict], root: str) -> nx.DiGraph:
    known = {root} | {p['id'] for p in packages} | {t['id'] for t in tasks}
    graph = nx.DiGraph(); graph.add_nodes_from(known)
    for p in packages:
        if p['parent'] not in known: raise g.Blocked('unknown parent %r for %s' % (p['parent'], p['id']))
        graph.add_edge(p['parent'], p['id'])
    for t in tasks:
        if t['parent'] not in known: raise g.Blocked('unknown parent %r for %s' % (t['parent'], t['id']))
        graph.add_edge(t['parent'], t['id'])
    if not nx.is_arborescence(graph): raise g.Blocked('the breakdown is not a tree')
    return graph


def rollup(tasks: list[dict], *, drop_null_parents: bool) -> dict:
    frame = pd.DataFrame(tasks)
    grouped = frame.groupby('parent', dropna=drop_null_parents)['days'].sum()
    return {('<none>' if pd.isna(k) else k): int(v) for k, v in grouped.items()}

# ---------------------------------------------------------------- P144
def variance_join(planned: list[dict], observed: list[dict], *, how: str) -> pd.DataFrame:
    left = pd.DataFrame(planned); right = pd.DataFrame(observed)
    joined = left.merge(right, on='task', how=how, validate='one_to_one', indicator=True)
    if how == 'outer' and not (joined['_merge'] == 'both').all():
        missing = joined.loc[joined['_merge'] != 'both', 'task'].tolist()
        raise g.Blocked('tasks without a matching record on both sides: %s' % missing)
    return joined

# ---------------------------------------------------------------- P145
def daily_load(assignments: list[dict]) -> pd.DataFrame:
    frame = pd.DataFrame(assignments)
    frame['days'] = [list(a['days']) for a in assignments]
    return frame.explode('days')


def exact_share(hours: str, days: int) -> Fraction:
    if days <= 0: raise g.Blocked('effort declared against %d scheduled days' % days)
    return Fraction(hours) / days

# ---------------------------------------------------------------- P146
def raci_matrix(cells: list[dict], tasks: list[str], roles: list[str]) -> pd.DataFrame:
    frame = pd.DataFrame(cells)
    wide = frame.pivot(index='task', columns='role', values='code')
    return wide.reindex(index=tasks, columns=roles)


def cell_states(matrix: pd.DataFrame) -> dict:
    return {(t, r): ('assigned' if isinstance(v, str) else 'unassigned')
            for t, row in matrix.iterrows() for r, v in row.items()}

# ---------------------------------------------------------------- P147
def cost_totals(rows: list[tuple], *, category: str | None = None) -> dict:
    con = duckdb.connect()
    con.execute('CREATE TABLE ledger(id VARCHAR, category VARCHAR, amount DECIMAL(18,2))')
    con.executemany('INSERT INTO ledger VALUES (?,?,?)', rows)
    where = '' if category is None else ' WHERE category = ?'
    args = [] if category is None else [category]
    total = con.execute('SELECT sum(amount) FROM ledger' + where, args).fetchone()[0]
    typed = con.execute('SELECT sum(amount) FROM ledger' + where, args).description[0][1]
    mean = con.execute('SELECT avg(amount) FROM ledger' + where, args)
    return {'total': total, 'total_type': str(typed), 'mean': mean.fetchone()[0],
            'mean_type': str(mean.description[0][1])}


def stored_amounts(values: list) -> list:
    con = duckdb.connect(); con.execute('CREATE TABLE t(amount DECIMAL(18,2))')
    con.executemany('INSERT INTO t VALUES (?)', [(v,) for v in values])
    return [r[0] for r in con.execute('SELECT amount FROM t').fetchall()]


def two_places(value: str) -> Decimal:
    amount = Decimal(value)
    if -amount.as_tuple().exponent > 2: raise g.Blocked('more than two decimal places: %s' % value)
    return amount

# ---------------------------------------------------------------- P148
TEST_SOURCE = '''import pytest


def test_total_is_exact():
    assert 1 == 1


def test_cutoff_excludes_future():
    assert 2 == 2


def test_rejects_bad_precision():
    assert 1 == 2


@pytest.mark.skip(reason='desktop spreadsheet editing is not executed here')
def test_desktop_edit():
    assert False


@pytest.fixture
def broken_input():
    raise RuntimeError('fixture could not build the input')


def test_reads_workbook(broken_input):
    assert True
'''


def run_named_tests(source: str) -> bytes:
    directory = Path(tempfile.mkdtemp())
    (directory / 'test_contract.py').write_text(source)
    report = directory / 'report.xml'
    subprocess.run([sys.executable, '-m', 'pytest', '-q', '-p', 'no:cacheprovider',
                    '--junitxml', str(report), str(directory)], cwd=directory,
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, check=False)
    if not report.exists(): raise g.Blocked('pytest produced no report')
    return report.read_bytes()


def test_states(report: bytes, *, rule: str) -> dict:
    root = DET.fromstring(report.decode())
    suite = root[0] if root.tag == 'testsuites' else root
    out = {}
    for case in suite.findall('testcase'):
        if rule == 'absence_of_failure':
            out[case.get('name')] = 'passed' if not case.findall('failure') else 'failed'
        elif rule == 'declared':
            marks = [c.tag for c in case if c.tag in ('failure', 'error', 'skipped')]
            out[case.get('name')] = {'failure': 'failed', 'error': 'errored', 'skipped': 'skipped'}[marks[0]] if marks else 'passed'
        else:
            raise g.Blocked('unknown result rule %r' % rule)
    return out

# ---------------------------------------------------------------- P149
AGING_EDGES = [0, 30, 60, 90]


def aging_buckets(days: list[int], *, mode: str) -> list:
    series = pd.Series(days)
    if mode == 'right_closed':
        return [str(v) for v in pd.cut(series, bins=AGING_EDGES, labels=['1-30', '31-60', '61-90'])]
    if mode == 'covering':
        edges = [-np.inf, 0] + AGING_EDGES[1:] + [np.inf]
        labels = ['not_due', 'due_1_30', 'due_31_60', 'due_61_90', 'due_over_90']
        return [str(v) for v in pd.cut(series, bins=edges, labels=labels)]
    raise g.Blocked('unknown bucket mode %r' % mode)

# ---------------------------------------------------------------- P150
def closing_cash(opening: str, moves: dict, days: list[str], *, skipna: bool) -> list:
    series = pd.Series([moves.get(d, np.nan) for d in days], index=days, dtype='float64')
    return [None if pd.isna(v) else float(v) for v in float(opening) + series.cumsum(skipna=skipna)]


def exact_closing(opening: str, moves: dict, days: list[str]) -> list[Decimal]:
    steps = []
    for d in days:
        if d not in moves: raise g.Blocked('no declared movement for %s' % d)
        steps.append(Decimal(str(moves[d])))
    return list(itertools.accumulate(steps, initial=Decimal(opening)))[1:]


def weekly_days(days: list[str], *, method: str) -> dict:
    index = pd.to_datetime(days)
    if method == 'date_range':
        weeks = pd.date_range(days[0], days[-1], freq='W')
        return {str(w.date()): int(((index > w - pd.Timedelta(days=7)) & (index <= w)).sum()) for w in weeks}
    if method == 'resample':
        counts = pd.Series(1, index=index).resample('W').sum()
        return {str(k.date()): int(v) for k, v in counts.items()}
    raise g.Blocked('unknown weekly method %r' % method)


C = []
def case(i):
    def reg(fn): C.append((i, fn)); return fn
    return reg


@case(141)
def a_bar_takes_a_width_not_an_end_date():
    intervals = [{'id': 'v20', 'start': '2026-09-08', 'finish': '2026-09-10'},
                 {'id': 'ga', 'start': '2026-09-14', 'finish': '2026-09-14'}]
    right = drawn_widths(bar_spans(intervals, width_from='width'))
    g.equal(right, [2.0, 0.0])  # the second interval is a same-day milestone and draws nothing
    wrong = drawn_widths(bar_spans(intervals, width_from='finish'))
    g.equal(wrong[0], 20706.0)  # passing the finish as the width draws a bar 20706 days long, without complaint
    g.equal(wrong[0] > right[0] * 10000, True)
    day = datetime.date(2026, 9, 8)
    back = mdates.num2date(mdates.date2num(day))
    g.equal(isinstance(back, datetime.datetime), True)  # num2date returns a tz-aware datetime, not a date
    g.equal(str(back.tzinfo), 'UTC'); g.equal(back.date(), day)
    g.rejects(g.Blocked, lambda: bar_spans([{'id': 'x', 'start': '2026-09-10', 'finish': '2026-09-08'}], width_from='width'))
    columns = ['task', 'start', 'finish']
    rows = [{'task': 'v20', 'start': '2026-09-08', 'finish': '2026-09-10'}]
    silent = workbook_rows(schedule_table(rows, columns, declare_columns=False))
    g.equal(silent[0], ['Column1', 'Column2', 'Column3'])  # add_table overwrote the header row that was already written
    declared = workbook_rows(schedule_table(rows, columns, declare_columns=True))
    g.equal(declared[0], columns); g.equal(declared[1], ['v20', '2026-09-08', '2026-09-10'])
    return {'zero_duration_milestone_draws_nothing': True, 'end_date_as_width_reproduced': True,
            'wrong_width_days': wrong[0], 'right_width_days': right[0],
            'add_table_overwrites_the_header_row': True, 'declared_columns_keep_the_source_headers': True}


@case(142)
def working_day_counts_are_half_open_and_offsets_can_raise():
    g.equal(working_days('2026-09-07', '2026-09-11', inclusive=False), 3)  # Sep 7 is excluded by the calendar
    g.equal(working_days('2026-09-07', '2026-09-11', inclusive=True), 4)  # an inclusive finish needs the extra day
    g.equal(working_days('2026-09-08', '2026-09-08', inclusive=False), 0)  # a same-day span counts zero, not one
    g.equal(np.is_busday('2026-09-07', busdaycal=CALENDAR).item(), False)
    g.rejects(ValueError, lambda: offset_day('2026-09-07', 0, roll='raise'))  # the default roll raises on a nonworking start
    g.equal(offset_day('2026-09-07', 0, roll='forward'), '2026-09-08')
    g.equal(offset_day('2026-09-07', 0, roll='backward'), '2026-09-04')  # four calendar days apart from the forward answer
    tasks = [{'id': 't1', 'days': 2, 'after': []}, {'id': 't2', 'days': 3, 'after': ['t1']},
             {'id': 't3', 'days': 1, 'after': ['t1']}, {'id': 't9', 'days': 0, 'after': ['t2', 't3']}]
    plan = build_plan(tasks)
    g.equal(earliest_starts(plan), {'t1': 0, 't2': 2, 't3': 2, 't9': 5})
    forward = nx.DiGraph(); forward.add_edges_from([('t3', 't9'), ('t1', 't9'), ('t2', 't9')])
    backward = nx.DiGraph(); backward.add_edges_from([('t2', 't9'), ('t1', 't9'), ('t3', 't9')])
    g.equal(list(nx.topological_sort(forward)) == list(nx.topological_sort(backward)), False)  # insertion order decides it
    g.equal(list(nx.lexicographical_topological_sort(forward)), list(nx.lexicographical_topological_sort(backward)))
    g.rejects(g.Blocked, lambda: build_plan([{'id': 't1', 'days': 1, 'after': ['nope']}]))
    g.rejects(g.Blocked, lambda: build_plan([{'id': 'a', 'days': 1, 'after': ['b']}, {'id': 'b', 'days': 1, 'after': ['a']}]))
    return {'busday_count_is_half_open': True, 'inclusive_finish_needs_plus_one': True,
            'default_roll_raises_on_a_nonworking_start': True, 'forward_and_backward_rolls_differ': True,
            'topological_sort_depends_on_insertion_order': True, 'lexicographical_sort_is_stable': True}


@case(143)
def a_rollup_over_a_null_parent_loses_the_row():
    tasks = [{'id': 't1', 'parent': 'wp1', 'days': 2}, {'id': 't2', 'parent': 'wp1', 'days': 3},
             {'id': 't3', 'parent': None, 'days': 4}, {'id': 't4', 'parent': 'wp2', 'days': 1}]
    dropped = rollup(tasks, drop_null_parents=True)
    g.equal(dropped, {'wp1': 5, 'wp2': 1})
    g.equal(sum(dropped.values()), 6); g.equal(sum(t['days'] for t in tasks), 10)  # four task-days vanished, no error
    kept = rollup(tasks, drop_null_parents=False)
    g.equal(kept, {'wp1': 5, 'wp2': 1, '<none>': 4})
    g.equal(sum(kept.values()), sum(t['days'] for t in tasks))  # the rollup now reconciles to the source
    stray = nx.DiGraph(); stray.add_nodes_from(['t1', 't2']); stray.add_edge('wp_typo', 't2')
    g.equal(sorted(stray.nodes), ['t1', 't2', 'wp_typo'])  # add_edge invents a node for a mistyped parent
    packages = [{'id': 'wp1', 'parent': 'p'}, {'id': 'wp2', 'parent': 'p'}]
    tree = wbs_graph([{'id': 't1', 'parent': 'wp1'}, {'id': 't2', 'parent': 'wp2'}], packages, 'p')
    g.equal(sorted(nx.descendants(tree, 'p')), ['t1', 't2', 'wp1', 'wp2'])
    two_parents = nx.DiGraph([('p', 'wp1'), ('p', 'wp2'), ('wp1', 't1'), ('wp2', 't1')])
    g.equal(nx.is_directed_acyclic_graph(two_parents), True)  # a task under two parents is still a DAG
    g.equal(nx.is_arborescence(two_parents), False)  # only the tree check catches it
    g.equal(len(nx.descendants(two_parents, 'p')), 3)  # and descendants would count that task once against two parents
    g.rejects(g.Blocked, lambda: wbs_graph([{'id': 't1', 'parent': 'nope'}], packages, 'p'))
    g.rejects(g.Blocked, lambda: wbs_graph([{'id': 't1', 'parent': 'wp1'}, {'id': 't1', 'parent': 'wp2'}], packages, 'p'))
    return {'groupby_drops_null_keys_by_default': True, 'dropped_task_days': 4,
            'add_edge_invents_a_node_for_a_typo': True, 'is_dag_accepts_two_parents': True,
            'is_arborescence_rejects_them': True}


@case(144)
def a_validated_join_still_drops_what_it_cannot_match():
    planned = [{'task': 't1', 'planned': '2026-09-22'}, {'task': 't2', 'planned': '2026-09-18'},
               {'task': 't3', 'planned': '2026-09-24'}]
    observed = [{'task': 't1', 'actual': '2026-09-22'}, {'task': 't2', 'actual': None},
                {'task': 't9', 'actual': '2026-09-19'}]
    inner = variance_join(planned, observed, how='inner')
    g.equal(len(inner), 2)  # validate='one_to_one' checked uniqueness, not coverage
    g.equal(sorted(inner['task']), ['t1', 't2'])  # t3 and t9 are gone and nothing was raised
    g.equal(len(planned) + 1, 4)
    try:
        variance_join(planned, observed, how='outer'); missing = []
    except g.Blocked as exc:
        missing = sorted(str(exc).split(': ')[1].strip("[]").replace("'", '').split(', '))
    g.equal(missing, ['t3', 't9'])  # an outer join with an indicator names both sides that did not match
    kept = pd.DataFrame(observed)
    g.equal(kept.loc[kept['task'] == 't2', 'actual'].isna().item(), True)  # a missing actual date stays missing
    g.rejects(pd.errors.MergeError,
              lambda: variance_join(planned + [{'task': 't1', 'planned': '2026-09-30'}], observed, how='inner'))
    return {'validate_one_to_one_checks_uniqueness_only': True, 'inner_join_dropped_rows': 2,
            'outer_join_with_indicator_names_them': True, 'missing_actual_dates_retained': True,
            'a_duplicated_key_is_rejected': True}


@case(145)
def an_empty_day_list_becomes_a_phantom_day():
    assignments = [{'id': 'a1', 'resource': 'ops', 'hours': '16', 'days': ['2026-09-08', '2026-09-09']},
                   {'id': 'a2', 'resource': 'ops', 'hours': '0', 'days': []}]
    exploded = daily_load(assignments)
    g.equal(len(exploded), 3)  # two real days plus one row the empty list produced
    g.equal(int(exploded['days'].isna().sum()), 1)  # the milestone's day is null, not absent
    real = exploded.dropna(subset=['days'])
    g.equal(len(real), 2)
    g.equal(exact_share('16', 2), Fraction(8))  # exact arithmetic, so the daily shares add back to the declared effort
    g.equal(sum([exact_share('16', 2)] * 2), Fraction(16))
    sixths = [exact_share('8', 6)] * 6
    g.equal(sum(sixths), Fraction(8))  # the exact split always adds back to the declared effort
    g.equal(sum([8 / 6] * 6), 7.999999999999999)  # the float split of the same figures does not
    g.equal(sum([8 / 6] * 6) == 8.0, False)
    g.equal(sum([10 / 3] * 3) == 10.0, True)  # and a neighbouring figure does land back, so the drift is not predictable
    g.rejects(g.Blocked, lambda: exact_share('4', 0))  # effort declared against a zero-duration milestone is Blocked
    return {'explode_emits_a_null_row_for_an_empty_list': True, 'phantom_rows': 1,
            'exact_fraction_split_conserves_effort': True, 'float_split_does_not': True,
            'effort_on_a_zero_day_task_is_blocked': True}


@case(146)
def a_blank_matrix_cell_needs_an_explicit_state():
    cells = [{'task': 't1', 'role': 'ops', 'code': 'A'}, {'task': 't1', 'role': 'legal', 'code': 'R'},
             {'task': 't2', 'role': 'ops', 'code': 'R'}]
    tasks = ['t1', 't2']; roles = ['ops', 'legal']
    matrix = raci_matrix(cells, tasks, roles)
    g.equal(matrix.loc['t1', 'ops'], 'A'); g.equal(matrix.loc['t2', 'ops'], 'R')
    g.equal(pd.isna(matrix.loc['t2', 'legal']), True)  # reindex filled the unassigned cell with a null
    states = cell_states(matrix)
    g.equal(states[('t2', 'legal')], 'unassigned'); g.equal(states[('t1', 'legal')], 'assigned')
    g.equal(sum(v == 'assigned' for v in states.values()), 3)
    g.equal(sum(v == 'unassigned' for v in states.values()), 1)
    accountable = {t: [r for r in roles if matrix.loc[t, r] == 'A'] for t in tasks}
    g.equal(accountable['t1'], ['ops']); g.equal(accountable['t2'], [])  # t2 has no accountable role and the check can say so
    duplicate = cells + [{'task': 't1', 'role': 'ops', 'code': 'R'}]
    g.rejects(ValueError, lambda: raci_matrix(duplicate, tasks, roles))  # pivot refuses a duplicated task/role pair
    return {'reindex_fills_unassigned_cells_with_null': True, 'explicit_state_column_separates_them': True,
            'assigned_cells': 3, 'unassigned_cells': 1, 'a_task_without_an_accountable_role_is_visible': True,
            'pivot_rejects_duplicate_pairs': True}


@case(147)
def exactness_survives_the_sum_and_stops_at_the_division():
    rows = [('c1', 'travel', Decimal('1200.25')), ('c2', 'travel', Decimal('-130.00')),
            ('c3', 'training', Decimal('3485.25'))]
    totals = cost_totals(rows)
    g.equal(totals['total'], Decimal('4555.50'))  # the sum stays exact
    g.equal(isinstance(totals['total'], Decimal), True)
    g.equal(totals['total_type'], 'DECIMAL(38,2)')  # widened, still fixed point
    g.equal(totals['mean_type'], 'DOUBLE')  # but avg leaves exact arithmetic entirely
    g.equal(isinstance(totals['mean'], float), True)
    empty = cost_totals(rows, category='facilities')
    g.equal(empty['total'], None)  # a category with no transactions sums to NULL, not to zero
    g.equal(empty['total'] == Decimal('0.00'), False)
    g.equal(stored_amounts([1200.255, 0.1 + 0.2]), [Decimal('1200.26'), Decimal('0.30')])  # floats are rounded in silently
    g.equal(stored_amounts([Decimal('1.005')]), [Decimal('1.01')])  # so is an over-precise Decimal
    g.equal(two_places('1200.25'), Decimal('1200.25'))
    g.rejects(g.Blocked, lambda: two_places('1.005'))  # the rejection has to happen before the value reaches the column
    g.rejects(InvalidOperation, lambda: two_places('not a number'))
    return {'sum_of_decimal_is_exact': True, 'empty_category_sums_to_null': True,
            'avg_returns_double': True, 'floats_and_overprecise_decimals_are_rounded_silently': True,
            'precision_must_be_rejected_before_insert': True}


@case(148)
def a_pass_is_the_absence_of_a_child_element():
    report = run_named_tests(TEST_SOURCE)
    naive = test_states(report, rule='absence_of_failure')
    declared = test_states(report, rule='declared')
    g.equal(len(declared), 5)
    g.equal(declared['test_total_is_exact'], 'passed')
    g.equal(declared['test_rejects_bad_precision'], 'failed')
    g.equal(declared['test_desktop_edit'], 'skipped')  # a skip is a declared state, not a result
    g.equal(declared['test_reads_workbook'], 'errored')  # so is a fixture that could not build the input
    g.equal(sum(v == 'passed' for v in naive.values()), 4)  # the naive rule counts the skip and the error as passes
    g.equal(sum(v == 'passed' for v in declared.values()), 2)
    g.equal(naive['test_desktop_edit'], 'passed'); g.equal(naive['test_reads_workbook'], 'passed')
    root = DET.fromstring(report.decode())
    suite = root[0] if root.tag == 'testsuites' else root
    g.equal(suite.get('errors'), '1'); g.equal(suite.get('failures'), '1'); g.equal(suite.get('skipped'), '1')
    g.rejects(g.Blocked, lambda: test_states(report, rule='optimistic'))
    return {'tests_executed': 5, 'naive_rule_reports_passes': 4, 'declared_rule_reports_passes': 2,
            'a_passing_testcase_has_no_child_element': True, 'skip_and_error_are_not_passes': True}


@case(149)
def a_right_closed_bucket_drops_the_invoice_due_today():
    days = [0, 1, 30, 31, 60, 61, 90, 91, 120]
    naive = aging_buckets(days, mode='right_closed')
    g.equal(naive[0], 'nan')  # due today falls outside (0, 30]
    g.equal(naive[-1], 'nan'); g.equal(naive[-2], 'nan')  # and so does anything past the last edge
    g.equal(sum(v == 'nan' for v in naive), 3)
    g.equal(naive[1], '1-30'); g.equal(naive[2], '1-30'); g.equal(naive[3], '31-60')
    covering = aging_buckets(days, mode='covering')
    g.equal(sum(v == 'nan' for v in covering), 0)  # every invoice lands in exactly one bucket
    g.equal(covering, ['not_due', 'due_1_30', 'due_1_30', 'due_31_60', 'due_31_60',
                       'due_61_90', 'due_61_90', 'due_over_90', 'due_over_90'])
    amounts = [Decimal('100.00')] * len(days)
    per_bucket = {}
    for bucket, amount in zip(covering, amounts): per_bucket[bucket] = per_bucket.get(bucket, Decimal('0.00')) + amount
    g.equal(sum(per_bucket.values()), Decimal('900.00'))  # the buckets reconcile to the outstanding total
    naive_total = sum(a for b, a in zip(naive, amounts) if b != 'nan')
    g.equal(naive_total, Decimal('600.00'))  # the right-closed buckets do not
    g.rejects(g.Blocked, lambda: aging_buckets(days, mode='loose'))
    return {'pd_cut_is_right_closed': True, 'due_today_and_over_90_fall_outside': True,
            'unbucketed_invoices': 3, 'covering_bins_reconcile_to_the_total': True,
            'right_closed_total': '600.00', 'covering_total': '900.00'}


@case(150)
def a_missing_day_does_not_stop_a_running_total():
    days = ['2026-09-21', '2026-09-22', '2026-09-23']
    moves = {'2026-09-21': 100.0, '2026-09-23': -50.0}  # nothing declared for the 22nd
    skipped = closing_cash('2500', moves, days, skipna=True)
    g.equal(skipped, [2600.0, None, 2550.0])  # the running total steps over the gap and looks correct after it
    strict = closing_cash('2500', moves, days, skipna=False)
    g.equal(strict, [2600.0, None, None])  # propagating the gap makes the missing day visible downstream
    complete = dict(moves, **{'2026-09-22': 0.0})
    g.equal(closing_cash('2500', complete, days, skipna=True), [2600.0, 2600.0, 2550.0])
    g.equal(exact_closing('2500', complete, days), [Decimal('2600.0'), Decimal('2600.0'), Decimal('2550.0')])
    g.rejects(g.Blocked, lambda: exact_closing('2500', moves, days))  # the exact path refuses an undeclared day
    span = [str(d.date()) for d in pd.date_range('2026-09-07', '2026-09-30')]
    g.equal(len(span), 24)
    ranged = weekly_days(span, method='date_range')
    g.equal(sum(ranged.values()), 21)  # the W anchor drops the days after the last Sunday
    g.equal(sorted(ranged), ['2026-09-13', '2026-09-20', '2026-09-27'])
    resampled = weekly_days(span, method='resample')
    g.equal(sum(resampled.values()), 24)  # resample keeps the partial week and the day count reconciles
    g.equal(sorted(resampled)[-1], '2026-10-04')
    g.rejects(g.Blocked, lambda: weekly_days(span, method='monthly'))
    return {'cumsum_skips_nulls_by_default': True, 'skipna_false_propagates_the_gap': True,
            'date_range_W_drops_the_partial_last_week': True, 'dropped_days': 3,
            'resample_reconciles_to_the_day_count': True, 'exact_path_refuses_an_undeclared_day': True}


def run() -> dict:
    g.ASSERTIONS = 0; rows = []
    for i, fn in C:
        before = g.ASSERTIONS
        try: r = fn(); state = 'passed'; error = None
        except Exception: r = {}; state = 'failed'; error = traceback.format_exc()
        rows.append({'proof': i, 'test': fn.__name__, 'status': state, 'assertions': g.ASSERTIONS - before, 'result': r, 'error': error})
    packages = ['duckdb', 'matplotlib', 'networkx', 'numpy', 'openpyxl', 'pandas', 'XlsxWriter', 'defusedxml', 'pytest']
    here = Path(__file__).resolve().parent
    return {'scope': 'Bounded authored schedule, cost and cash fixtures and reproduced library defaults; the original P141-P150 chains and their artifacts were not rerun.',
            'test_cases': len(rows), 'passed': sum(r['status'] == 'passed' for r in rows), 'failed': sum(r['status'] == 'failed' for r in rows),
            'assertions': g.ASSERTIONS, 'full_chains_executed': 0, 'network_calls': 0, 'business_actions': 0,
            'source_sha256': hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
            'dependency_sha256': {name: hashlib.sha256((here / name).read_bytes()).hexdigest() for name in ['handoff_guards_v1.py']},
            'library_versions': {n: importlib.metadata.version(n) for n in packages}, 'results': rows}


if __name__ == '__main__':
    p = argparse.ArgumentParser(); p.add_argument('--report', type=Path, required=True)
    args = p.parse_args()
    result = run(); args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(json.dumps(result, indent=2, sort_keys=True, allow_nan=False) + '\n')
    print(json.dumps({k: result[k] for k in ['test_cases', 'passed', 'failed', 'assertions']}))
    if result['failed']:
        print('\n'.join(r['error'] for r in result['results'] if r['error'])); raise SystemExit(1)
