"""Fourth bounded adapter suite: P31-P40. Requires handoff_guards_v1.py, handoff_guards_v2.py and handoff_guards_v3.py
from this TXT. Every comparison, join, pivot, grounding, conversion, serialization, query, build, and execution below is a
library primitive call; local code declares fixtures and policies and turns primitive results into Blocked outcomes.
All runtime inputs are authored fixtures. No network, model client, native Office application, or business action.
Run: python handoff_guards_v4.py --report report.json
"""
from __future__ import annotations
import argparse, hashlib, html as html_text, importlib.metadata, io, json, logging, os, tempfile, traceback, warnings
from pathlib import Path
from typing import Literal
from urllib.parse import quote
import clingo
import docx
import duckdb
import kuzu
import markdown
import mkdocs.commands.build
import mkdocs.config
import networkx as nx
import pandas as pd
import plotly.graph_objects as go
import plotly.io as pio
import pm4py
from csv_diff import compare as csv_compare
from deepdiff import DeepDiff
from docx.table import Table
from lxml import html
from pm4py.objects.bpmn.importer.variants import lxml as bpmn_lxml
from pm4py.objects.petri_net.obj import Marking, PetriNet
from pm4py.objects.petri_net.utils import petri_utils
from pydantic import BaseModel, ConfigDict, Field, ValidationError
from rdflib import Graph, Literal as RdfLiteral, Namespace, URIRef
from rdflib.compare import isomorphic
from rdflib.namespace import XSD
from rdflib.term import Node
import handoff_guards_v1 as g
import handoff_guards_v3 as g3
warnings.simplefilter('ignore'); logging.disable(logging.WARNING)

# ---------------------------------------------------------------- P31
class AttendanceRow(BaseModel):
    model_config = ConfigDict(strict=True, extra='forbid')
    meeting: str = Field(min_length=1)
    present: int = Field(ge=0)
    threshold: int = Field(ge=1)
    majority_reachable: bool


def attendance_chart(rows: list[dict], *, include_plotlyjs) -> dict:
    typed = [AttendanceRow.model_validate(r) for r in rows]
    for r in typed:
        if r.majority_reachable != (r.present >= r.threshold): raise g.Blocked('majority flag disagrees with present and threshold for ' + r.meeting)
    x = [r.meeting for r in typed]
    fig = go.Figure([go.Bar(name='present', x=x, y=[r.present for r in typed]),
                     go.Scatter(name='threshold', x=x, y=[r.threshold for r in typed], mode='lines+markers')])
    data = json.loads(pio.to_json(fig, validate=True))['data']
    diff = DeepDiff({'present': [r.present for r in typed], 'threshold': [r.threshold for r in typed]},
                    {'present': data[0]['y'], 'threshold': data[1]['y']}, zip_ordered_iterables=True)
    if diff: raise g.Blocked('chart payload differs from the rows: ' + diff.to_json())
    page = pio.to_html(fig, include_plotlyjs=include_plotlyjs, full_html=True, div_id='attendance', validate=True)
    external = g3.external_assets(page)
    if external: raise g.Blocked('external asset dependency: ' + external[0])
    return {'traces': [t['name'] for t in data], 'meetings': len(typed), 'bytes': len(page.encode())}

# ---------------------------------------------------------------- P32
def coverage_matrix(tags: list[dict], *, actions: list[str], meetings: list[str]) -> dict:
    df = pd.DataFrame(tags)
    if not df['meeting'].isin(meetings).all() or not df['action'].isin(actions).all(): raise g.Blocked('tag outside the declared axes')
    if set(df['state']) - {'accepted', 'tie'}: raise g.Blocked('unknown tag state')
    accepted = df[df['state'] == 'accepted']
    matrix = accepted.assign(n=1).pivot_table(index='action', columns='meeting', values='n', aggfunc='sum', fill_value=0)
    matrix = matrix.reindex(index=actions, columns=meetings, fill_value=0)
    if int(matrix.to_numpy().sum()) != len(accepted): raise g.Blocked('matrix does not conserve the accepted tags')
    return {'meetings': list(meetings), 'matrix': {a: [int(v) for v in matrix.loc[a]] for a in actions},
            'accepted': int(len(accepted)), 'tied_excluded': int((df['state'] == 'tie').sum())}

# ---------------------------------------------------------------- P33
CLOSED_WORLD = 'never(E) :- step(E), not discussed(E). #show never/1.'
THREE_VALUED = """never(E) :- step(E), not discussed(E), not tie(E), coverage(complete).
possibly(E) :- step(E), tie(E).
unknown(E) :- step(E), not discussed(E), not tie(E), not coverage(complete).
#show never/1. #show possibly/1. #show unknown/1."""


def ground_steps(steps: list[dict], tags: list[dict], *, coverage: str, program: str) -> dict:
    ctl = clingo.Control(['0', '--warn=no-atom-undefined'])
    with ctl.backend() as be:
        for s in steps: be.add_rule([be.add_atom(clingo.Function('step', [clingo.String(s['id'])]))])
        for t in tags:
            if t['state'] not in {'accepted', 'tie'}: raise g.Blocked('unknown tag state')
            name = 'discussed' if t['state'] == 'accepted' else 'tie'
            be.add_rule([be.add_atom(clingo.Function(name, [clingo.String(t['step'])]))])
        if coverage == 'complete': be.add_rule([be.add_atom(clingo.Function('coverage', [clingo.Function('complete')]))])
    ctl.add('base', [], program); ctl.ground([('base', [])])
    shown = []; ctl.solve(on_model=lambda m: shown.extend(m.symbols(shown=True)))
    out = {}
    for s in shown: out.setdefault(s.name, []).append(s.arguments[0].string)
    return {k: sorted(v) for k, v in sorted(out.items())}


def undiscussed_steps(steps: list[dict], tags: list[dict], *, coverage: str) -> dict:
    ids = [s['id'] for s in steps]
    if pd.Index(ids).has_duplicates: raise g.Blocked('duplicate step identity')
    result = ground_steps(steps, tags, coverage=coverage, program=THREE_VALUED)
    return {'never': result.get('never', []), 'possibly': result.get('possibly', []), 'unknown': result.get('unknown', []),
            'lemmas': {s['id']: s['lemma'] for s in steps}, 'coverage': coverage}

# ---------------------------------------------------------------- P34
def build_net(name: str, places: list[str], transitions: list[tuple[str, str]], arcs: list[tuple[str, str]], source: str, sink: str):
    net = PetriNet(name); P = {p: PetriNet.Place(p) for p in places}; T = {n: PetriNet.Transition(n, label) for n, label in transitions}
    for p in P.values(): net.places.add(p)
    for t in T.values(): net.transitions.add(t)
    for a, b in arcs: petri_utils.add_arc_from_to(P.get(a) or T[a], T.get(b) or P[b], net)
    return net, Marking({P[source]: 1}), Marking({P[sink]: 1})


def xor_net(): return build_net('xor', ['source', 'p1', 'p2', 'sink'], [('ta', 'A'), ('tb', 'B'), ('tc', 'C'), ('td', 'D')],
                                [('source', 'ta'), ('source', 'tb'), ('ta', 'p1'), ('tb', 'p2'), ('p1', 'tc'), ('p2', 'td'), ('tc', 'sink'), ('td', 'sink')], 'source', 'sink')


def and_xor_net(): return build_net('and_xor', ['source', 'q1', 'q2', 'sink'], [('t0', 'split'), ('t1', 'A'), ('t2', 'B')],
                                    [('source', 't0'), ('t0', 'q1'), ('t0', 'q2'), ('q1', 't1'), ('q2', 't2'), ('t1', 'sink'), ('t2', 'sink')], 'source', 'sink')


def sequence_net(): return build_net('sequence', ['source', 'p1', 'sink'], [('ta', 'A'), ('tb', 'B')],
                                     [('source', 'ta'), ('ta', 'p1'), ('p1', 'tb'), ('tb', 'sink')], 'source', 'sink')


def structural_decision(net, im, fm) -> dict:
    workflow = bool(pm4py.check_is_workflow_net(net))
    # marked graph: every place has one producing and one consuming transition; the source has none before it and the sink none after it
    marked = all(len(p.in_arcs) == (0 if p in im else 1) and len(p.out_arcs) == (0 if p in fm else 1) for p in net.places)
    acyclic = bool(nx.is_directed_acyclic_graph(nx.DiGraph([(a.source.name, a.target.name) for a in net.arcs])))
    return {'workflow_net': workflow, 'marked_graph': marked, 'acyclic': acyclic}


def revised_rule(net, im, fm) -> bool:
    # The recorded revision: True only for an acyclic marked-graph workflow net, False otherwise.
    s = structural_decision(net, im, fm); return s['workflow_net'] and s['marked_graph'] and s['acyclic']


def soundness_state(net, im, fm) -> dict:
    s = structural_decision(net, im, fm)
    if not s['workflow_net']: return {'state': 'not_a_workflow_net', 'method': 'structural', **s}
    if s['marked_graph'] and s['acyclic']: return {'state': 'sound', 'method': 'structural: acyclic marked-graph workflow net', **s}
    sound, _ = pm4py.check_soundness(net, im, fm)
    return {'state': 'sound' if sound else 'unsound', 'method': 'woflan after structural undecided', **s}

# ---------------------------------------------------------------- P35
def gateway_contract(xml_text: str) -> nx.DiGraph:
    graph = g3.sequence_flow_graph(xml_text)
    offenders = sorted(n for n, d in graph.nodes(data=True) if not d['kind'].endswith('Gateway') and (graph.in_degree(n) > 1 or graph.out_degree(n) > 1))
    if offenders: raise g.Blocked('uncontrolled flow without a gateway at ' + ','.join(offenders) + ': split and join semantics undeclared')
    return graph


def petri_view(xml_text: str) -> dict:
    net, im, fm = pm4py.convert_to_petri_net(bpmn_lxml.import_from_string(xml_text))
    places = {p.name: (len(p.in_arcs), len(p.out_arcs)) for p in net.places}
    transitions = {t.label: (len(t.in_arcs), len(t.out_arcs)) for t in net.transitions if t.label}
    return {'places': places, 'transitions': transitions, 'soundness': soundness_state(net, im, fm)}

# ---------------------------------------------------------------- P36
def ordered_events(events: list[dict], *, lifecycle: str = 'complete', tiebreak: str | None = None) -> pd.DataFrame:
    df = pd.DataFrame(events)
    if df[['case_id', 'event_id']].duplicated().any(): raise g.Blocked('duplicate event identity')
    df['timestamp'] = pd.to_datetime(df['timestamp'], utc=True, errors='raise')
    df = pm4py.format_dataframe(df, case_id='case_id', activity_key='activity', timestamp_key='timestamp')
    kept = pm4py.filter_event_attribute_values(df, 'lifecycle', [lifecycle], level='event', retain=True)
    if kept.duplicated([g3.CASE, g3.TS], keep=False).any() and tiebreak is None: raise g.Blocked('timestamp ties need an ordering contract')
    return kept.sort_values([g3.CASE, g3.TS] + ([tiebreak] if tiebreak else []), kind='stable').reset_index(drop=True)


def checked_variants(events: list[dict], *, tiebreak: str | None = None) -> dict:
    kept = ordered_events(events, tiebreak=tiebreak)
    variants = pm4py.get_variants(kept, activity_key=g3.ACT, timestamp_key=g3.TS, case_id_key=g3.CASE)
    counted = {k: (v if isinstance(v, int) else len(v)) for k, v in variants.items()}
    if sum(counted.values()) != int(kept[g3.CASE].nunique()): raise g.Blocked('variant counts do not conserve the cases')
    return {'variants': [{'trace': list(k), 'cases': int(v)} for k, v in sorted(counted.items())], 'cases': int(sum(counted.values()))}

# ---------------------------------------------------------------- P37
FACTS = Namespace('http://example.org/facts/')
COUNT_QUERY = 'SELECT ?s WHERE { ?s <http://example.org/facts/count> ?v FILTER(?v = ?x) }'


def fact_graph(facts: list[dict]) -> tuple[Graph, str]:
    graph = Graph()
    for f in facts:
        subject = FACTS[quote(f['id'], safe='')]
        graph.add((subject, FACTS['lemma'], RdfLiteral(f['lemma'])))
        graph.add((subject, FACTS['count'], RdfLiteral(int(f['count']), datatype=XSD.integer)))
    turtle = graph.serialize(format='turtle')
    if not isomorphic(graph, Graph().parse(data=turtle, format='turtle')): raise g.Blocked('turtle round trip is not isomorphic')
    return graph, turtle


def facts_with_count(graph: Graph, value: int) -> list[str]:
    return sorted(str(r.s) for r in graph.query(COUNT_QUERY, initBindings={'x': RdfLiteral(int(value), datatype=XSD.integer)}))

# ---------------------------------------------------------------- P38
class EventNode(BaseModel):
    model_config = ConfigDict(strict=True, extra='forbid')
    id: str = Field(min_length=1)
    lemma: str = Field(min_length=1)

SCHEMA = ['CREATE NODE TABLE Event(id STRING, lemma STRING, PRIMARY KEY(id))', 'CREATE REL TABLE precedes(FROM Event TO Event)']


def graph_store(events: list[dict], edges: list[tuple[str, str]]) -> dict:
    typed = [EventNode.model_validate(e) for e in events]
    conn = kuzu.Connection(kuzu.Database(':memory:'))
    for ddl in SCHEMA: conn.execute(ddl)
    for e in typed: conn.execute('CREATE (:Event {id: $id, lemma: $lemma})', parameters={'id': e.id, 'lemma': e.lemma})
    for a, b in edges:
        created = conn.execute('MATCH (a:Event {id: $a}), (b:Event {id: $b}) CREATE (a)-[:precedes]->(b) RETURN count(*)', parameters={'a': a, 'b': b}).get_all()[0][0]
        if created != 1: raise g.Blocked('edge %s->%s references an unknown event' % (a, b))
    rows = conn.execute('MATCH (e:Event) RETURN e.id, e.lemma ORDER BY e.id').get_all()
    if DeepDiff(sorted([[e.id, e.lemma] for e in typed]), rows, zip_ordered_iterables=True): raise g.Blocked('graph readback differs from the rows')
    edge_count = conn.execute('MATCH (:Event)-[r:precedes]->(:Event) RETURN count(r)').get_all()[0][0]
    return {'events': len(rows), 'precedes': int(edge_count)}

# ---------------------------------------------------------------- P39
def table_markdown(rows: list[dict], columns: list[str]) -> str:
    cell = lambda v: html_text.escape(str(v)).replace('|', '\\|')
    lines = ['| ' + ' | '.join(columns) + ' |', '|' + '|'.join(' --- ' for _ in columns) + '|']
    lines += ['| ' + ' | '.join(cell(r[c]) for c in columns) + ' |' for r in rows]
    return '\n'.join(lines) + '\n'


def rendered_rows(html_page: str, table_xpath: str) -> list[list[str]]:
    return [[c.text_content() for c in r.xpath('./td')] for r in html.fromstring(html_page).xpath(table_xpath + '//tbody/tr')]


def build_site(md_text: str, *, highlightjs: bool) -> str:
    root = tempfile.mkdtemp(); os.makedirs(root + '/docs')
    Path(root, 'docs', 'index.md').write_text(md_text)
    Path(root, 'mkdocs.yml').write_text('site_name: proofs\nmarkdown_extensions: [tables]\ntheme:\n  name: mkdocs\n  highlightjs: %s\n' % str(highlightjs).lower())
    mkdocs.commands.build.build(mkdocs.config.load_config(config_file=root + '/mkdocs.yml', site_dir=root + '/site'))
    return Path(root, 'site', 'index.html').read_text()


def checked_site(rows: list[dict], columns: list[str]) -> dict:
    page = build_site(table_markdown(rows, columns), highlightjs=False)
    external = g3.external_assets(page)
    if external: raise g.Blocked('external asset dependency: ' + external[0])
    back = rendered_rows(page, '//*[@role="main"]//table[1]')
    diff = DeepDiff([[str(r[c]) for c in columns] for r in rows], back, zip_ordered_iterables=True)
    if diff: raise g.Blocked('site table differs from the rows: ' + diff.to_json())
    return {'rows': len(back), 'external_assets': 0}

# ---------------------------------------------------------------- P40
class StepRow(BaseModel):
    model_config = ConfigDict(strict=True, extra='forbid')
    id: str = Field(min_length=1)
    lemma: str = Field(min_length=1)
    quote: str = Field(min_length=1)

STEP_COLUMNS = list(StepRow.model_fields)


def step_document(steps: list[dict]) -> bytes:
    typed = [StepRow.model_validate(s) for s in steps]
    if pd.Index([s.id for s in typed]).has_duplicates: raise g.Blocked('duplicate step identity')
    d = docx.Document(); d.add_paragraph('ordered steps'); t = d.add_table(rows=1, cols=len(STEP_COLUMNS))
    for cell, name in zip(t.rows[0].cells, STEP_COLUMNS): cell.text = name
    for s in typed:
        for cell, name in zip(t.add_row().cells, STEP_COLUMNS): cell.text = getattr(s, name)
    buf = io.BytesIO(); d.save(buf); return buf.getvalue()


def read_step_document(data: bytes) -> list[dict]:
    d = docx.Document(io.BytesIO(data))
    tables = [b for b in d.iter_inner_content() if isinstance(b, Table)]
    if len(tables) != 1: raise g.Blocked('expected exactly one step table in body order')
    rows = [StepRow.model_validate(dict(zip(STEP_COLUMNS, [c.text for c in r.cells]))).model_dump() for r in tables[0].rows[1:]]
    if pd.Index([r['id'] for r in rows]).has_duplicates: raise g.Blocked('duplicate step identity on readback')
    return rows

C = []
def case(i):
    def reg(fn): C.append((i, fn)); return fn
    return reg


@case(31)
def threshold_bound_to_data():
    rows = [{'meeting': 'tsc-2023-11-08', 'present': 6, 'threshold': 9, 'majority_reachable': False},
            {'meeting': 'tsc-2023-12-06', 'present': 10, 'threshold': 9, 'majority_reachable': True},
            {'meeting': 'tsc-2025-02-05', 'present': 6, 'threshold': 10, 'majority_reachable': False}]
    fig = go.Figure(go.Bar(x=[r['meeting'] for r in rows], y=[r['present'] for r in rows])); fig.add_hline(y=10)
    payload = json.loads(pio.to_json(fig))
    g.equal([s['y0'] for s in payload['layout']['shapes']], [10]); g.equal(len(payload['data']), 1)  # the line is a layout shape, not data
    out = attendance_chart(rows, include_plotlyjs=True)
    g.equal(out['traces'], ['present', 'threshold'])
    g.rejects(g.Blocked, lambda: attendance_chart([dict(rows[1], majority_reachable=False)], include_plotlyjs=True))
    g.rejects(g.Blocked, lambda: attendance_chart(rows, include_plotlyjs='cdn'))
    g.rejects(ValidationError, lambda: attendance_chart([dict(rows[0], threshold='10')], include_plotlyjs=True))
    return {'hline_is_layout_shape_reproduced': True, 'per_meeting_threshold_in_data': True, 'flag_checked_against_data': True,
            'cdn_page': 'blocked', 'original_chain_rerun': False}


@case(32)
def coverage_matrix_declared_axes():
    tags = [{'meeting': 'm1', 'action': 'approve', 'state': 'accepted'}, {'meeting': 'm1', 'action': 'open', 'state': 'accepted'},
            {'meeting': 'm3', 'action': 'open', 'state': 'accepted'}, {'meeting': 'm3', 'action': 'approve', 'state': 'tie'}]
    df = pd.DataFrame(tags); con = duckdb.connect(); con.register('tags', df)
    pivoted = con.execute('PIVOT tags ON meeting USING count(*) GROUP BY action ORDER BY action').df()
    g.equal([c for c in pivoted.columns if c != 'action'], ['m1', 'm3'])  # the absent meeting has no column
    g.equal(list(df.assign(n=1).pivot_table(index='action', columns='meeting', values='n', aggfunc='sum', fill_value=0).columns), ['m1', 'm3'])
    out = coverage_matrix(tags, actions=['approve', 'open', 'remedy'], meetings=['m1', 'm2', 'm3'])
    g.equal(out['matrix'], {'approve': [1, 0, 0], 'open': [1, 0, 1], 'remedy': [0, 0, 0]}); g.equal(out['tied_excluded'], 1)
    g.rejects(g.Blocked, lambda: coverage_matrix(tags + [{'meeting': 'm9', 'action': 'open', 'state': 'accepted'}], actions=['approve', 'open', 'remedy'], meetings=['m1', 'm2', 'm3']))
    g.rejects(g.Blocked, lambda: coverage_matrix(tags + [{'meeting': 'm1', 'action': 'open', 'state': 'maybe'}], actions=['approve', 'open', 'remedy'], meetings=['m1', 'm2', 'm3']))
    return {'absent_column_dropped_reproduced': 'duckdb PIVOT and pandas pivot_table', 'declared_axes_reindexed': True,
            'ties_not_coverage': True, 'conservation_checked': True, 'original_chain_rerun': False}


@case(33)
def negation_as_failure_needs_coverage():
    steps = [{'id': 'e1', 'lemma': 'share'}, {'id': 'e2', 'lemma': 'ensure'}, {'id': 'e3', 'lemma': 'meeting'}, {'id': 'e4', 'lemma': 'meeting'}]
    tags = [{'step': 'e1', 'state': 'accepted'}, {'step': 'e2', 'state': 'tie'}]
    g.equal(ground_steps(steps, tags, coverage='unknown', program=CLOSED_WORLD), {'never': ['e2', 'e3', 'e4']})  # the tie-only step counts as never discussed
    complete = undiscussed_steps(steps, tags, coverage='complete')
    g.equal((complete['never'], complete['possibly'], complete['unknown']), (['e3', 'e4'], ['e2'], []))
    g.equal([complete['lemmas'][i] for i in complete['never']], ['meeting', 'meeting'])
    partial = undiscussed_steps(steps, tags, coverage='unknown')
    g.equal((partial['never'], partial['possibly'], partial['unknown']), ([], ['e2'], ['e3', 'e4']))
    g.rejects(g.Blocked, lambda: undiscussed_steps(steps + [steps[0]], tags, coverage='complete'))
    g.rejects(g.Blocked, lambda: undiscussed_steps(steps, tags + [{'step': 'e3', 'state': 'maybe'}], coverage='complete'))
    return {'closed_world_never_for_tie_only_step_reproduced': True, 'three_valued_states': ['never', 'possibly', 'unknown'],
            'coverage_fact_required_for_never': True, 'duplicate_lemma_steps_kept_distinct': True, 'original_chain_rerun': False}


@case(34)
def soundness_state_not_conflated():
    xor = xor_net(); and_xor = and_xor_net(); seq = sequence_net()
    g.equal(revised_rule(*xor), False); g.equal(soundness_state(*xor)['state'], 'sound')  # the revision's False is an undecided net that Woflan proves sound
    g.equal(soundness_state(*xor)['method'], 'woflan after structural undecided')
    g.equal(soundness_state(*and_xor)['state'], 'unsound'); g.equal(revised_rule(*and_xor), False)
    g.equal(soundness_state(*seq), {'state': 'sound', 'method': 'structural: acyclic marked-graph workflow net', 'workflow_net': True, 'marked_graph': True, 'acyclic': True})
    broken = build_net('broken', ['source', 'orphan', 'sink'], [('ta', 'A')], [('source', 'ta'), ('ta', 'sink')], 'source', 'sink')
    g.equal(soundness_state(*broken)['state'], 'not_a_workflow_net')
    return {'undecided_recorded_as_false_reproduced': True, 'states': ['sound', 'unsound', 'not_a_workflow_net'], 'method_named': True,
            'woflan_time_budget': 'not applied; fixture nets finish immediately', 'original_chain_rerun': False}


@case(35)
def engine_semantics_of_uncontrolled_flow():
    merge = petri_view(g3.BPMN_MERGE)
    g.equal(merge['places']['exi_start'], (1, 2)); g.equal(merge['places']['ent_c'], (2, 1)); g.equal(merge['transitions']['c'], (1, 1))
    g.equal(merge['soundness']['state'], 'sound')  # one branch runs and c fires once in the Petri view
    g.equal(g3.run_witness(g3.BPMN_MERGE, 'proc'), ['a', 'b', 'c', 'c'])  # both branches run and c fires twice in SpiffWorkflow
    g.rejects(g.Blocked, lambda: gateway_contract(g3.BPMN_MERGE))
    join = petri_view(g3.BPMN_JOIN)
    g.equal(join['transitions']['c'], (2, 1)); g.equal(join['soundness']['state'], 'sound')
    g.equal(g3.run_witness(g3.BPMN_JOIN, 'proc'), ['a', 'b', 'c'])
    g.equal(sorted(n for n, d in gateway_contract(g3.BPMN_JOIN).nodes(data=True) if d['kind'] != 'sequenceFlow'), ['a', 'b', 'c', 'end', 'join', 'split', 'start'])
    return {'petri_conversion_exclusive_for_uncontrolled_flow': True, 'spiffworkflow_parallel_for_uncontrolled_flow': True,
            'engines_disagree_yet_both_report_success': True, 'gateway_contract_blocks_uncontrolled_flow': True, 'original_chain_rerun': False}


@case(36)
def variants_as_tuples_with_ordering_contract():
    t0, t1, t2 = '2011-01-01T00:00:00Z', '2011-01-01T00:00:10Z', '2011-01-01T00:00:20Z'
    rows = [{'case_id': 'c1', 'event_id': 'e1', 'activity': 'A -> B', 'timestamp': t0, 'lifecycle': 'complete'},
            {'case_id': 'c1', 'event_id': 'e2', 'activity': 'C', 'timestamp': t1, 'lifecycle': 'complete'},
            {'case_id': 'c2', 'event_id': 'e3', 'activity': 'A', 'timestamp': t0, 'lifecycle': 'complete'},
            {'case_id': 'c2', 'event_id': 'e4', 'activity': 'B -> C', 'timestamp': t1, 'lifecycle': 'complete'}]
    out = checked_variants(rows)
    g.equal(out['variants'], [{'trace': ['A', 'B -> C'], 'cases': 1}, {'trace': ['A -> B', 'C'], 'cases': 1}]); g.equal(out['cases'], 2)
    g.equal({' -> '.join(v['trace']) for v in out['variants']}, {'A -> B -> C'})  # the joined string collapses two variants
    tied = [{'case_id': 'c3', 'event_id': 'e5', 'activity': 'X', 'timestamp': t0, 'lifecycle': 'complete'},
            {'case_id': 'c3', 'event_id': 'e6', 'activity': 'Y', 'timestamp': t1, 'lifecycle': 'complete'},
            {'case_id': 'c3', 'event_id': 'e7', 'activity': 'Z', 'timestamp': t1, 'lifecycle': 'complete'}]
    g.rejects(g.Blocked, lambda: checked_variants(tied))
    g.equal(checked_variants(tied, tiebreak='event_id'), checked_variants([tied[0], tied[2], tied[1]], tiebreak='event_id'))
    g.rejects(g.Blocked, lambda: checked_variants(rows + [rows[0]]))
    return {'joined_string_collision_reproduced': True, 'tuple_keys_kept': True, 'ties_need_contract': True, 'case_conservation': True, 'original_chain_rerun': False}


@case(37)
def iri_minting_and_typed_query():
    g.equal(isinstance(FACTS.count, Node), False); g.equal(isinstance(FACTS['count'], Node), True)  # attribute access returns str.count
    raw = Graph(); raw.add((URIRef('http://example.org/facts/a b'), FACTS['lemma'], RdfLiteral('x')))
    g.rejects(Exception, lambda: raw.serialize(format='turtle'))  # an IRI with a space cannot be serialized
    facts = [{'id': 'usc5-552-doj:u0001:s00#e7', 'lemma': 'make', 'count': 1}, {'id': 'a b', 'lemma': 'x', 'count': 2}, {'id': 'e9', 'lemma': 'state', 'count': 1}]
    graph, turtle = fact_graph(facts)
    g.equal('usc5-552-doj%3Au0001%3As00%23e7' in turtle, True); g.equal(len(graph), 6)
    g.equal(facts_with_count(graph, 1), ['http://example.org/facts/e9', 'http://example.org/facts/usc5-552-doj%3Au0001%3As00%23e7'])
    plain = Graph(); plain.add((FACTS['p1'], FACTS['count'], RdfLiteral('1'))); plain.add((FACTS['p2'], FACTS['count'], RdfLiteral(1)))
    g.equal([str(r.s) for r in plain.query('SELECT ?s WHERE { ?s <http://example.org/facts/count> ?v FILTER(?v = 1) }')], ['http://example.org/facts/p2'])
    hostile = '1) } UNION { ?s ?p ?o } FILTER(true'
    g.rejects(Exception, lambda: list(plain.query('SELECT ?s WHERE { ?s <http://example.org/facts/count> ?v FILTER(?v = ' + hostile + ') }')))
    g.equal(list(plain.query(COUNT_QUERY, initBindings={'x': RdfLiteral(hostile)})), [])
    return {'namespace_attribute_trap_reproduced': True, 'unencoded_iri_serialization_failure_reproduced': True, 'percent_encoded_round_trip_isomorphic': True,
            'plain_versus_typed_literal_reproduced': True, 'query_text_fixed_values_bound': True, 'original_chain_rerun': False}


@case(38)
def parameterized_graph_load():
    hostile = "x'}) RETURN 1; //"
    events = [{'id': 'e1', 'lemma': hostile}, {'id': 'e2', 'lemma': 'make'}]
    g.equal(graph_store(events, [('e1', 'e2')]), {'events': 2, 'precedes': 1})
    conn = kuzu.Connection(kuzu.Database(':memory:'))
    for ddl in SCHEMA: conn.execute(ddl)
    conn.execute("CREATE (:Event {id: 'e3', lemma: '" + hostile + "'})")
    g.equal(conn.execute('MATCH (e:Event) RETURN e.lemma').get_all(), [['x']])  # interpolation stored a different value without an error
    conn.execute('CREATE (:Event {id: $id, lemma: $lemma})', parameters={'id': 'e4', 'lemma': hostile})
    g.equal(conn.execute('MATCH (e:Event {id: $id}) RETURN e.lemma', parameters={'id': 'e4'}).get_all(), [[hostile]])
    silent = conn.execute('MATCH (a:Event {id: $a}), (b:Event {id: $b}) CREATE (a)-[:precedes]->(b) RETURN count(*)', parameters={'a': 'e3', 'b': 'zz'}).get_all()
    g.equal(silent, [[0]])  # an edge to an unknown event vanishes without an error
    g.rejects(g.Blocked, lambda: graph_store(events, [('e1', 'zz')]))
    g.rejects(ValidationError, lambda: graph_store(events + [{'id': 1, 'lemma': 'x'}], []))
    return {'cypher_interpolation_changed_value_reproduced': True, 'parameters_keep_value': True, 'lost_edge_reproduced_and_blocked': True,
            'typed_rows_before_load': True, 'results_ordered_explicitly': True, 'original_chain_rerun': False}


@case(39)
def site_tables_need_escaping_and_local_assets():
    raw = '| id | text |\n| --- | --- |\n| s2 | a | b |\n| s3 | <b>bold</b> |\n'
    g.equal(rendered_rows(markdown.markdown(raw, extensions=['tables']), '//table'), [['s2', 'a'], ['s3', 'bold']])  # a pipe loses a cell; HTML is rendered
    default_page = build_site(raw, highlightjs=True)
    g.equal([u.startswith('https://cdnjs.cloudflare.com/') for u in g3.external_assets(default_page)], [True, True, True])
    rows = [{'id': 's2', 'text': 'a | b'}, {'id': 's3', 'text': '<b>bold</b>'}, {'id': 's4', 'text': 'line1<br>line2'}]
    g.equal(checked_site(rows, ['id', 'text']), {'rows': 3, 'external_assets': 0})
    g.equal(len(rendered_rows(build_site(raw, highlightjs=False), '//table')) > 2, True)  # theme tables exist beyond the content table
    return {'pipe_cell_loss_reproduced': True, 'inline_html_rendered_reproduced': True, 'cdn_assets_in_default_theme_reproduced': True,
            'escaped_rows_round_trip': True, 'content_table_selected_by_role': True, 'original_chain_rerun': False}


@case(40)
def document_readback_keyed_by_identity():
    steps = [{'id': 'u0020:s00#e25', 'lemma': 'meeting', 'quote': 'the meeting'}, {'id': 'u0040:s01#e8', 'lemma': 'meeting', 'quote': 'a meeting'}]
    data = step_document(steps)
    d = docx.Document(io.BytesIO(data))
    g.equal([p.text for p in d.paragraphs], ['ordered steps']); g.equal([type(b).__name__ for b in d.iter_inner_content()], ['Paragraph', 'Table'])
    by_text = {s['lemma']: s for s in steps}
    g.equal(len(by_text), 1); g.equal(csv_compare(by_text, {'meeting': steps[1]})['changed'], [])  # keying by text loses a row and reports no change
    g.equal(bool(DeepDiff(steps, read_step_document(data), zip_ordered_iterables=True)), False)
    d.tables[0].rows[2].cells[0].text = steps[0]['id']; buf = io.BytesIO(); d.save(buf)
    g.rejects(g.Blocked, lambda: read_step_document(buf.getvalue()))
    g.rejects(g.Blocked, lambda: step_document(steps + [steps[0]]))
    return {'paragraphs_omit_table_text_reproduced': True, 'text_keyed_collapse_reproduced': True, 'identity_keyed_positional_readback': True,
            'duplicate_identity_blocked': True, 'original_chain_rerun': False}


def run() -> dict:
    g.ASSERTIONS = 0; rows = []
    for i, fn in C:
        before = g.ASSERTIONS
        try: r = fn(); state = 'passed'; error = None
        except Exception: r = {}; state = 'failed'; error = traceback.format_exc()
        rows.append({'proof': i, 'test': fn.__name__, 'status': state, 'assertions': g.ASSERTIONS - before, 'result': r, 'error': error})
    packages = ['pydantic', 'pandas', 'duckdb', 'deepdiff', 'pm4py', 'plotly', 'lxml', 'python-docx', 'csv-diff', 'SpiffWorkflow', 'networkx', 'clingo', 'rdflib', 'kuzu', 'mkdocs', 'Markdown']
    here = Path(__file__).resolve().parent
    return {'scope': 'Bounded authored handoff fixtures and reproduced library defaults; the original P31-P40 chains and their artifacts were not rerun.',
            'test_cases': len(rows), 'passed': sum(r['status'] == 'passed' for r in rows), 'failed': sum(r['status'] == 'failed' for r in rows),
            'assertions': g.ASSERTIONS, 'full_chains_executed': 0, 'network_calls': 0, 'business_actions': 0,
            'source_sha256': hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
            'dependency_sha256': {name: hashlib.sha256((here / name).read_bytes()).hexdigest() for name in ['handoff_guards_v1.py', 'handoff_guards_v2.py', 'handoff_guards_v3.py']},
            'library_versions': {n: importlib.metadata.version(n) for n in packages}, 'results': rows}


if __name__ == '__main__':
    p = argparse.ArgumentParser(); p.add_argument('--report', type=Path, required=True); args = p.parse_args()
    result = run(); args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(json.dumps(result, indent=2, sort_keys=True, allow_nan=False) + '\n')
    print(json.dumps({k: result[k] for k in ['test_cases', 'passed', 'failed', 'assertions', 'full_chains_executed']}))
    if result['failed']:
        print('\n'.join(r['error'] for r in result['results'] if r['error'])); raise SystemExit(1)
