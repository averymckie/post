"""Third bounded adapter suite: P21-P30. Requires handoff_guards_v1.py and handoff_guards_v2.py from this TXT.
Every comparison, join, interval test, difference, similarity, grounding, and execution below is a library
primitive call; local code declares fixtures and policies and turns primitive results into Blocked outcomes.
Runtime inputs are authored fixtures, except the P26 token, which reproduces the bytes recorded in the P26 output.
No network, model client, native Office application, or business action.
Run: NLTK_DATA=<directory holding corpora/wordnet.zip> python handoff_guards_v3.py --report report.json
"""
from __future__ import annotations
import argparse, base64, hashlib, importlib.metadata, io, json, traceback, warnings
from collections import Counter
from datetime import date
from pathlib import Path
from typing import Literal
import clingo
import docx
import duckdb
import ftfy, ftfy.badness
import networkx as nx
import nltk
import numpy as np
import openpyxl
import pandas as pd
import plotly.graph_objects as go
import plotly.io as pio
import pm4py
import portion as P
from csv_diff import compare as csv_compare
from deepdiff import DeepDiff
from docx.table import Table
from jinja2 import Environment, StrictUndefined, UndefinedError
from lxml import etree, html
from pydantic import BaseModel, ConfigDict, Field, ValidationError, model_validator
from python_calamine import CalamineWorkbook
from sklearn.feature_extraction.text import CountVectorizer, TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity
from SpiffWorkflow.bpmn.parser.BpmnParser import BpmnParser
from SpiffWorkflow.bpmn.workflow import BpmnWorkflow
from SpiffWorkflow.util.task import TaskState
import handoff_guards_v1 as g
import handoff_guards_v2 as g2
warnings.simplefilter('ignore')


def records(frame: pd.DataFrame) -> list[dict]:
    # NaN/NaT never enter the evidence; missing aggregates are None.
    return frame.astype(object).where(frame.notna(), None).to_dict('records')

# ---------------------------------------------------------------- P21
class CaseRow(BaseModel):
    model_config = ConfigDict(strict=True, extra='forbid')
    case_id: str = Field(min_length=1)
    deadline: str | None
    enddate: str | None


def deadline_outcomes(rows: list[dict], *, cutoff: str, boundary: Literal['closed', 'open']) -> dict:
    if boundary not in {'closed', 'open'}: raise g.Blocked('deadline boundary policy must be declared')
    df = pd.DataFrame([CaseRow.model_validate(r).model_dump() for r in rows])
    if (df.groupby('case_id')[['deadline', 'enddate']].nunique(dropna=False) > 1).any().any():
        raise g.Blocked('case attributes disagree across rows')
    one = df.drop_duplicates('case_id')
    con = duckdb.connect(); con.register('one', one)
    typed = con.execute('select case_id, deadline::TIMESTAMPTZ as deadline, enddate::TIMESTAMPTZ as enddate, '
                        '?::TIMESTAMPTZ as cutoff from one order by case_id', [cutoff]).df()
    out = {}
    for r in typed.itertuples():
        if pd.isna(r.deadline): out[r.case_id] = 'unknown_deadline'; continue
        window = P.closed(-P.inf, r.deadline) if boundary == 'closed' else P.open(-P.inf, r.deadline)
        if pd.isna(r.enddate): out[r.case_id] = 'open_within_deadline' if r.cutoff in window else 'open_overdue_at_cutoff'
        else: out[r.case_id] = 'by_deadline' if r.enddate in window else 'after_deadline'
    return out

# ---------------------------------------------------------------- P22
APPROVED_LABELS = {'department': {'General', 'Experts', 'Customer contact'},
                   'channel': {'Desk', 'Internet', 'Intern', 'Post', 'e-mail'}}
GROUP_SQL = """select {key} as label, count(*) as cases,
 count(*) filter (where closed) as closed_cases, count(*) filter (where not closed) as open_cases,
 avg(epoch(last_ts::TIMESTAMPTZ - first_ts::TIMESTAMPTZ)) filter (where closed) / 86400.0 as mean_days_closed,
 count(*) filter (where outcome = 'after_deadline') as after_deadline,
 count(*) filter (where outcome = 'unknown_deadline') as unknown_deadline
 from cases group by 1 order by 1 nulls first"""


def group_measures(cases: list[dict], *, key: str) -> list[dict]:
    if key not in APPROVED_LABELS: raise g.Blocked('group key must be a declared attribute')
    df = pd.DataFrame(cases)
    if df['case_id'].duplicated().any(): raise g.Blocked('duplicate case')
    unapproved = sorted(set(df[key].dropna()) - APPROVED_LABELS[key])
    if unapproved: raise g.Blocked('unapproved group label (no silent normalization): ' + repr(unapproved))
    con = duckdb.connect(); con.register('cases', df)
    out = con.execute(GROUP_SQL.format(key=key)).df()
    if int(out['cases'].sum()) != len(df): raise g.Blocked('group sizes do not conserve the case population')
    return records(out)

# ---------------------------------------------------------------- P23
CASE, ACT, TS = 'case:concept:name', 'concept:name', 'time:timestamp'


def raw_dfg(events: list[dict]) -> dict:
    df = pd.DataFrame(events); df['timestamp'] = pd.to_datetime(df['timestamp'], utc=True)
    df = pm4py.format_dataframe(df, case_id='case_id', activity_key='activity', timestamp_key='timestamp')
    dfg, _, _ = pm4py.discover_dfg(df)
    return {a + '->' + b: n for (a, b), n in dfg.items()}


def checked_dfg(events: list[dict], *, lifecycle: str = 'complete', tiebreak: str | None = None) -> dict:
    df = pd.DataFrame(events)
    if not {'case_id', 'event_id', 'activity', 'timestamp', 'lifecycle'} <= set(df): raise g.Blocked('missing event fields')
    if df[['case_id', 'event_id']].duplicated().any(): raise g.Blocked('duplicate event identity')
    df['timestamp'] = pd.to_datetime(df['timestamp'], utc=True, errors='raise')
    df = pm4py.format_dataframe(df, case_id='case_id', activity_key='activity', timestamp_key='timestamp')
    kept = pm4py.filter_event_attribute_values(df, 'lifecycle', [lifecycle], level='event', retain=True)
    if kept.duplicated([CASE, TS], keep=False).any() and tiebreak is None:
        raise g.Blocked('timestamp ties need an ordering contract')
    kept = kept.sort_values([CASE, TS] + ([tiebreak] if tiebreak else []), kind='stable').reset_index(drop=True)
    dfg, starts, ends = pm4py.discover_dfg(kept, activity_key=ACT, timestamp_key=TS, case_id_key=CASE)
    following = kept.groupby(CASE)[ACT].shift(-1)
    declared = {k: v for k, v in Counter(zip(kept[ACT], following)).items() if isinstance(k[1], str)}
    if DeepDiff(dict(dfg), declared, zip_ordered_iterables=True): raise g.Blocked('engine order differs from the declared order')
    return {'edges': {a + '->' + b: n for (a, b), n in sorted(dfg.items())}, 'start': dict(starts), 'end': dict(ends),
            'events_used': int(len(kept)), 'events_excluded_by_lifecycle': int(len(df) - len(kept)),
            'cases': int(kept[CASE].nunique()), 'basis': 'observation'}

# ---------------------------------------------------------------- P24
class ChartRow(BaseModel):
    model_config = ConfigDict(strict=True, extra='forbid')
    activity: str = Field(min_length=1)
    mean_hours: float | None
    events: int = Field(ge=0)


def external_assets(page: str) -> list[str]:
    doc = html.fromstring(page)
    return [u for u in doc.xpath('//script/@src') + doc.xpath('//link/@href') if u.startswith(('http:', 'https:', '//'))]


def chart_page(rows: list[dict], *, include_plotlyjs) -> dict:
    typed = [ChartRow.model_validate(r) for r in rows]
    source = {'x': [r.activity for r in typed], 'y': [r.mean_hours for r in typed], 'customdata': [[r.events] for r in typed]}
    fig = go.Figure(go.Bar(**source, hovertemplate='%{x}: %{y} h over %{customdata[0]} events<extra></extra>'))
    payload = json.loads(pio.to_json(fig, validate=True))['data'][0]
    diff = DeepDiff(source, {k: payload[k] for k in source}, zip_ordered_iterables=True)
    if diff: raise g.Blocked('chart payload differs from the source rows: ' + diff.to_json())
    page = pio.to_html(fig, include_plotlyjs=include_plotlyjs, full_html=True, div_id='chart', validate=True)
    external = external_assets(page)
    if external: raise g.Blocked('external asset dependency: ' + external[0])
    return {'bytes': len(page.encode()), 'bars': len(typed), 'missing_values_kept': sum(r.mean_hours is None for r in typed)}

# ---------------------------------------------------------------- P25
class RequiredActionRow(BaseModel):
    model_config = ConfigDict(strict=True, extra='forbid')
    fact_id: str = Field(min_length=1)
    actor: str = Field(min_length=1)
    modality: Literal['obligatory', 'permitted', 'prohibited']
    action: str = Field(min_length=1)
    condition: str = Field(min_length=1)
    quote: str = Field(min_length=1)
    source_sentence: str = Field(min_length=1)

    @model_validator(mode='after')
    def quote_inside_source(self):
        if self.quote not in self.source_sentence: raise ValueError('quote is not a substring of its source sentence')
        return self

ACTION_COLUMNS = list(RequiredActionRow.model_fields)
RENDERING = Environment(undefined=StrictUndefined).from_string('{{ actor }} {{ modality }} {{ action }} when {{ condition }}')


def action_document(rows: list[dict]) -> bytes:
    typed = [RequiredActionRow.model_validate(r) for r in rows]
    d = docx.Document(); t = d.add_table(rows=1, cols=len(ACTION_COLUMNS) + 1)
    for cell, name in zip(t.rows[0].cells, ACTION_COLUMNS + ['rendering']): cell.text = name
    for r in typed:
        cells = t.add_row().cells
        for cell, name in zip(cells, ACTION_COLUMNS): cell.text = getattr(r, name)
        cells[-1].text = RENDERING.render(**r.model_dump())
    buf = io.BytesIO(); d.save(buf); return buf.getvalue()


def read_action_document(data: bytes) -> list[dict]:
    d = docx.Document(io.BytesIO(data))
    tables = [b for b in d.iter_inner_content() if isinstance(b, Table)]
    if len(tables) != 1: raise g.Blocked('expected exactly one action table in body order')
    if [c.text for c in tables[0].rows[0].cells] != ACTION_COLUMNS + ['rendering']: raise g.Blocked('column contract changed')
    out = []
    for row in tables[0].rows[1:]:
        values = [c.text for c in row.cells]
        typed = RequiredActionRow.model_validate(dict(zip(ACTION_COLUMNS, values[:-1])))
        if RENDERING.render(**typed.model_dump()) != values[-1]: raise g.Blocked('rendering does not match its typed fields')
        out.append(typed.model_dump())
    return out

# ---------------------------------------------------------------- P26
RECORDED_P26_TOKEN = 'It’s'.encode('utf-8').decode('latin-1').encode('utf-8').decode('latin-1')
ACTIONS = {'open': 'open the tsc meeting', 'record': 'record the tsc meeting', 'close': 'close the tsc meeting',
           'schedule': 'schedule the tsc meeting', 'chair': 'chair the tsc meeting', 'announce': 'announce the tsc meeting',
           'direct': 'direct the tsc meeting', 'remedy': 'remedy conflict'}


def checked_text(text: str) -> str:
    if ftfy.badness.is_bad(text): raise g.Blocked('mojibake detected; decode the source bytes with the right codec before tokenizing')
    return text


def tag_candidates(sentence: str, actions: dict[str, str], *, margin: float = 0.05) -> dict:
    checked_text(sentence)
    names = list(actions); docs = [actions[n] for n in names]
    binary = CountVectorizer(token_pattern=r'[a-z]+', binary=True).fit(docs + [sentence])
    overlap = np.asarray(binary.transform(docs).multiply(binary.transform([sentence])).sum(axis=1)).ravel()
    weighted = TfidfVectorizer(token_pattern=r'[a-z]+').fit(docs + [sentence])
    similarity = cosine_similarity(weighted.transform(docs), weighted.transform([sentence])).ravel()
    ranked = sorted(zip(names, similarity.tolist(), overlap.tolist()), key=lambda t: -t[1])
    state = 'tie' if ranked[0][1] - ranked[1][1] < margin else 'candidate'
    return {'state': state, 'overlap_winners': sorted(n for n, o in zip(names, overlap) if o == overlap.max()),
            'ranked': [{'action': n, 'tfidf_cosine': round(s, 4), 'shared_tokens': int(o)} for n, s, o in ranked],
            'authorization': 'none: a candidate needs an approved scoped lexicon and polarity check (P10)'}

# ---------------------------------------------------------------- P27
CONFORMANCE_SQL = """select department, split, count(*) as cases,
 count(*) filter (where state in ('fit', 'deviates')) as assessed, count(*) filter (where state = 'fit') as fit,
 count(*) filter (where state = 'not_run') as not_run,
 avg(fitness) filter (where state in ('fit', 'deviates')) as mean_fitness
 from joined group by 1, 2 order by 1 nulls first, 2"""


def conformance_by_group(record_rows: list[dict], conformance: list[dict], partition: list[dict]) -> dict:
    rec, conf, part = (pd.DataFrame(x) for x in (record_rows, conformance, partition))
    j = rec.merge(conf, on='case_id', how='outer', validate='one_to_one', indicator=True)
    if not j['_merge'].eq('both').all():
        raise g.Blocked('conformance coverage mismatch: ' + repr(sorted(j.loc[j['_merge'] != 'both', 'case_id'])))
    j = j.drop(columns='_merge').merge(part, on='case_id', how='outer', validate='one_to_one', indicator=True)
    if not j['_merge'].eq('both').all(): raise g.Blocked('partition label coverage mismatch')
    if set(j['state']) - {'fit', 'deviates', 'not_run', 'unknown'}: raise g.Blocked('unrecognized replay state')
    con = duckdb.connect(); con.register('joined', j.drop(columns='_merge'))
    return {'groups': records(con.execute(CONFORMANCE_SQL).df()),
            'in_sample_rows_present': bool((j['split'] == 'train').any()), 'regulatory_compliance_established': False}

# ---------------------------------------------------------------- P28
MEETING_PROGRAM = """meeting(M) :- attendance(M, _).
discussed(M, N) :- meeting(M), N = #count{ A : tag(M, A, accepted) }.
tied(M, N) :- meeting(M), N = #count{ A : tag(M, A, tie) }.
untagged(M) :- meeting(M), not tag(M, _, _).
unmatched(M) :- tag(M, _, _), not meeting(M).
:- attendance(M, true), attendance(M, false).
#show discussed/2. #show tied/2. #show untagged/1. #show unmatched/1."""


def meeting_join(decisions: list[dict], tags: list[dict]) -> dict:
    ctl = clingo.Control(['0'])
    with ctl.backend() as be:
        for d in decisions:
            flag = clingo.Function('true' if d['majority_reachable'] else 'false')
            be.add_rule([be.add_atom(clingo.Function('attendance', [clingo.String(d['meeting']), flag]))])
        for t in tags:
            if t['state'] not in {'accepted', 'tie'}: raise g.Blocked('unknown tag state')
            be.add_rule([be.add_atom(clingo.Function('tag', [clingo.String(t['meeting']), clingo.String(t['action']), clingo.Function(t['state'])]))])
    ctl.add('base', [], MEETING_PROGRAM); ctl.ground([('base', [])])
    shown = []; result = ctl.solve(on_model=lambda m: shown.extend(m.symbols(shown=True)))
    if not result.satisfiable: raise g.Blocked('conflicting attendance rows for one meeting')
    out = {'discussed': {}, 'tied': {}, 'untagged': [], 'unmatched': []}
    for s in shown:
        if s.name in ('discussed', 'tied'): out[s.name][s.arguments[0].string] = s.arguments[1].number
        else: out[s.name].append(s.arguments[0].string)
    out['untagged'].sort(); out['unmatched'].sort()
    return out

# ---------------------------------------------------------------- P29
BPMN_NS = {'bpmn': 'http://www.omg.org/spec/BPMN/20100524/MODEL'}
BPMN_HEAD = '<?xml version="1.0" encoding="UTF-8"?>\n<bpmn:definitions xmlns:bpmn="http://www.omg.org/spec/BPMN/20100524/MODEL" id="defs" targetNamespace="http://example.org/p29">\n<bpmn:process id="proc" isExecutable="true">\n'
BPMN_TAIL = '</bpmn:process>\n</bpmn:definitions>\n'
BPMN_MERGE = BPMN_HEAD + '''<bpmn:startEvent id="start"><bpmn:outgoing>f1</bpmn:outgoing><bpmn:outgoing>f2</bpmn:outgoing></bpmn:startEvent>
<bpmn:manualTask id="a" name="a"><bpmn:incoming>f1</bpmn:incoming><bpmn:outgoing>f3</bpmn:outgoing></bpmn:manualTask>
<bpmn:manualTask id="b" name="b"><bpmn:incoming>f2</bpmn:incoming><bpmn:outgoing>f4</bpmn:outgoing></bpmn:manualTask>
<bpmn:manualTask id="c" name="c"><bpmn:incoming>f3</bpmn:incoming><bpmn:incoming>f4</bpmn:incoming><bpmn:outgoing>f5</bpmn:outgoing></bpmn:manualTask>
<bpmn:endEvent id="end"><bpmn:incoming>f5</bpmn:incoming></bpmn:endEvent>
<bpmn:sequenceFlow id="f1" sourceRef="start" targetRef="a"/><bpmn:sequenceFlow id="f2" sourceRef="start" targetRef="b"/>
<bpmn:sequenceFlow id="f3" sourceRef="a" targetRef="c"/><bpmn:sequenceFlow id="f4" sourceRef="b" targetRef="c"/>
<bpmn:sequenceFlow id="f5" sourceRef="c" targetRef="end"/>
''' + BPMN_TAIL
BPMN_JOIN = BPMN_HEAD + '''<bpmn:startEvent id="start"><bpmn:outgoing>f0</bpmn:outgoing></bpmn:startEvent>
<bpmn:parallelGateway id="split"><bpmn:incoming>f0</bpmn:incoming><bpmn:outgoing>f1</bpmn:outgoing><bpmn:outgoing>f2</bpmn:outgoing></bpmn:parallelGateway>
<bpmn:manualTask id="a" name="a"><bpmn:incoming>f1</bpmn:incoming><bpmn:outgoing>f3</bpmn:outgoing></bpmn:manualTask>
<bpmn:manualTask id="b" name="b"><bpmn:incoming>f2</bpmn:incoming><bpmn:outgoing>f4</bpmn:outgoing></bpmn:manualTask>
<bpmn:parallelGateway id="join"><bpmn:incoming>f3</bpmn:incoming><bpmn:incoming>f4</bpmn:incoming><bpmn:outgoing>f5</bpmn:outgoing></bpmn:parallelGateway>
<bpmn:manualTask id="c" name="c"><bpmn:incoming>f5</bpmn:incoming><bpmn:outgoing>f6</bpmn:outgoing></bpmn:manualTask>
<bpmn:endEvent id="end"><bpmn:incoming>f6</bpmn:incoming></bpmn:endEvent>
<bpmn:sequenceFlow id="f0" sourceRef="start" targetRef="split"/><bpmn:sequenceFlow id="f1" sourceRef="split" targetRef="a"/>
<bpmn:sequenceFlow id="f2" sourceRef="split" targetRef="b"/><bpmn:sequenceFlow id="f3" sourceRef="a" targetRef="join"/>
<bpmn:sequenceFlow id="f4" sourceRef="b" targetRef="join"/><bpmn:sequenceFlow id="f5" sourceRef="join" targetRef="c"/>
<bpmn:sequenceFlow id="f6" sourceRef="c" targetRef="end"/>
''' + BPMN_TAIL


def sequence_flow_graph(xml_text: str) -> nx.DiGraph:
    root = etree.fromstring(xml_text.encode(), parser=etree.XMLParser(resolve_entities=False, no_network=True))
    graph = nx.DiGraph()
    for el in root.xpath('//bpmn:process/*[@id]', namespaces=BPMN_NS): graph.add_node(el.get('id'), kind=etree.QName(el).localname)
    for f in root.xpath('//bpmn:sequenceFlow', namespaces=BPMN_NS): graph.add_edge(f.get('sourceRef'), f.get('targetRef'), flow=f.get('id'))
    return graph


def run_witness(xml_text: str, process_id: str) -> list[str]:
    parser = BpmnParser(); parser.add_bpmn_xml(etree.fromstring(xml_text.encode()), filename='fixture.bpmn')
    wf = BpmnWorkflow(parser.get_spec(process_id)); wf.do_engine_steps(); trace = []
    while not wf.is_completed():
        ready = sorted(wf.get_tasks(state=TaskState.READY), key=lambda t: t.task_spec.name)
        if not ready or len(trace) > 100: raise g.Blocked('workflow stalled or ran away')
        ready[0].run(); trace.append(ready[0].task_spec.name); wf.do_engine_steps()
    return trace


def execution_checked(xml_text: str, process_id: str, forced_edges: list[tuple[str, str]]) -> dict:
    graph = sequence_flow_graph(xml_text)
    for a, b in forced_edges:
        if a not in graph or b not in graph or not nx.has_path(graph, a, b): raise g.Blocked('forced edge %s->%s has no sequence-flow path' % (a, b))
    trace = run_witness(xml_text, process_id)
    repeated = {k: v for k, v in Counter(trace).items() if v != 1}
    if repeated: raise g.Blocked('tasks executed more than once (uncontrolled merge): ' + repr(repeated))
    return {'witness_trace': trace, 'structurally_forced_edges': len(forced_edges)}

# ---------------------------------------------------------------- P30
def workbook_roundtrip(rows: list[dict], *, profile: Literal['typed', 'json_text']) -> dict:
    if profile not in {'typed', 'json_text'}: raise g.Blocked('cell profile must be declared')
    columns = list(rows[0]); wb = openpyxl.Workbook(); ws = wb.active; ws.append(columns)
    for r in rows:
        if list(r) != columns: raise g.Blocked('ragged rows')
        ws.append([json.dumps(r[c], allow_nan=False) if profile == 'json_text' else r[c] for c in columns])
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    sheet = CalamineWorkbook.from_filelike(buf).get_sheet_by_index(0).to_python(skip_empty_area=False)
    if sheet[0] != columns: raise g.Blocked('header changed')
    back = [dict(zip(columns, [g.strict_json_loads(v) if profile == 'json_text' else v for v in row])) for row in sheet[1:]]
    diff = DeepDiff(rows, back, zip_ordered_iterables=True)
    return {'profile': profile, 'rows_back': len(back), 'diff': json.loads(diff.to_json()) if diff else {}}


def checked_roundtrip(rows: list[dict], *, profile: str) -> dict:
    result = workbook_roundtrip(rows, profile=profile)
    if result['diff']: raise g.Blocked('typed readback differs: ' + json.dumps(result['diff'], sort_keys=True))
    return result

# ---------------------------------------------------------------- glue register: library primitives for v1/v2 helpers
class OccurrenceReceipt(BaseModel):
    model_config = ConfigDict(strict=True, extra='forbid')
    kind: Literal['observed_completion']
    case_id: str = Field(min_length=1)
    action_id: str = Field(min_length=1)
    receipt_id: str = Field(min_length=1)
    timestamp: str = Field(min_length=1)


def process_map_diff(expected: dict, rows: list[dict]) -> dict:
    if pd.Index([r['id'] for r in rows]).has_duplicates: raise g.Blocked('duplicate BPMN identity')
    diff = DeepDiff(expected, {r['id']: r['kind'] for r in rows}, zip_ordered_iterables=True)
    if diff: raise g.Blocked('process map differs: ' + diff.to_json())
    return {'retained': len(rows)}


def rosters_at(meeting: date, rosters: list[dict]) -> list[dict]:
    return [r for r in rosters if meeting in P.closedopen(r['valid_from'], r['valid_to'])]


def verdict(fn, *args, **kwargs) -> str:
    try: fn(*args, **kwargs); return 'accepted'
    except (g.Blocked, ValidationError): return 'rejected'

C = []
def case(i):
    def reg(fn): C.append((i, fn)); return fn
    return reg


@case(21)
def deadline_instants_and_censoring():
    a, b = '2011-10-30 02:30:00+02:00', '2011-10-30 02:00:00+01:00'  # a is 00:30Z, b is 01:00Z
    g.equal(a <= b, False)  # byte order of the strings says a is after b
    naive_le, naive_a = duckdb.execute('select ?::TIMESTAMP <= ?::TIMESTAMP, ?::TIMESTAMP', [a, b, a]).fetchone()
    g.equal(naive_le, False); g.equal(str(naive_a), '2011-10-30 02:30:00')  # the naive cast drops the offset
    g.equal(duckdb.execute('select ?::TIMESTAMPTZ <= ?::TIMESTAMPTZ', [a, b]).fetchone()[0], True)
    unknown, coalesced = duckdb.execute('select NULL::TIMESTAMPTZ <= ?::TIMESTAMPTZ, coalesce(NULL::TIMESTAMPTZ <= ?::TIMESTAMPTZ, false)', [b, b]).fetchone()
    g.equal(unknown, None); g.equal(coalesced, False)
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter('always'); mixed = pd.to_datetime(pd.Series([a, b]), utc=False)
    g.equal(str(mixed.dtype), 'object'); g.require(any(issubclass(w.category, FutureWarning) for w in caught), 'mixed offsets did not warn')
    g.equal(str(pd.to_datetime(pd.Series([a, b]), utc=True).dtype), 'datetime64[ns, UTC]')
    rows = [{'case_id': 'c1', 'deadline': b, 'enddate': a},
            {'case_id': 'c2', 'deadline': '2011-12-06 13:41:31+01:00', 'enddate': None},
            {'case_id': 'c3', 'deadline': None, 'enddate': a},
            {'case_id': 'c4', 'deadline': '2012-02-01 00:00:00+01:00', 'enddate': None},
            {'case_id': 'c5', 'deadline': b, 'enddate': b}]
    cutoff = '2012-01-23 00:00:00+00:00'
    g.equal(deadline_outcomes(rows, cutoff=cutoff, boundary='closed'),
            {'c1': 'by_deadline', 'c2': 'open_overdue_at_cutoff', 'c3': 'unknown_deadline', 'c4': 'open_within_deadline', 'c5': 'by_deadline'})
    g.equal(deadline_outcomes(rows, cutoff=cutoff, boundary='open')['c5'], 'after_deadline')
    g.rejects(g.Blocked, lambda: deadline_outcomes(rows + [{'case_id': 'c1', 'deadline': a, 'enddate': a}], cutoff=cutoff, boundary='closed'))
    g.rejects(ValidationError, lambda: deadline_outcomes([{'case_id': 'c9', 'deadline': b, 'enddate': a, 'extra': 1}], cutoff=cutoff, boundary='closed'))
    g.rejects(g.Blocked, lambda: deadline_outcomes(rows, cutoff=cutoff, boundary='either'))
    return {'string_and_naive_timestamp_order_reproduced_wrong': True, 'instant_comparison': 'TIMESTAMPTZ',
            'null_deadline_state': 'unknown_deadline', 'open_cases_split_by_cutoff': True, 'boundary_policy_explicit': True,
            'pandas_mixed_offsets_default': 'object dtype with FutureWarning', 'original_chain_rerun': False}


@case(22)
def group_measures_keep_missing_and_censored():
    cases = [{'case_id': 'a', 'department': 'General', 'first_ts': '2011-10-01 10:00:00+02:00', 'last_ts': '2011-10-03 10:00:00+02:00', 'closed': True, 'outcome': 'by_deadline'},
             {'case_id': 'b', 'department': None, 'first_ts': '2011-10-01 10:00:00+02:00', 'last_ts': None, 'closed': False, 'outcome': 'open_overdue_at_cutoff'},
             {'case_id': 'c', 'department': 'General', 'first_ts': '2011-10-01 10:00:00+02:00', 'last_ts': '2011-10-05 10:00:00+02:00', 'closed': True, 'outcome': 'after_deadline'},
             {'case_id': 'd', 'department': 'Experts', 'first_ts': '2011-10-01 10:00:00+02:00', 'last_ts': None, 'closed': False, 'outcome': 'unknown_deadline'}]
    df = pd.DataFrame(cases)
    g.equal(int(df.groupby('department').size().sum()), 3)  # pandas default drops the missing department
    g.equal(int(df.groupby('department', dropna=False).size().sum()), 4)
    out = group_measures(cases, key='department')
    g.equal([(r['label'], r['cases'], r['closed_cases'], r['open_cases'], r['after_deadline'], r['unknown_deadline']) for r in out],
            [(None, 1, 0, 1, 0, 0), ('Experts', 1, 0, 1, 0, 1), ('General', 2, 2, 0, 1, 0)])
    g.equal([r['mean_days_closed'] for r in out], [None, None, 3.0])
    g.equal(len(pd.DataFrame(cases + [dict(cases[0], case_id='e', department='General ')]).groupby('department').size()), 3)
    g.rejects(g.Blocked, lambda: group_measures(cases + [dict(cases[0], case_id='e', department='General ')], key='department'))
    g.rejects(g.Blocked, lambda: group_measures(cases + [cases[0]], key='department'))
    g.rejects(g.Blocked, lambda: group_measures(cases, key='resource'))
    return {'pandas_dropna_default_loss_reproduced': True, 'missing_group_retained': True, 'population_conserved': True,
            'label_variants_not_normalized': 'blocked', 'duration_only_for_closed_cases': True, 'original_chain_rerun': False}


@case(23)
def dfg_tie_and_lifecycle_contract():
    t0, t1, t2 = '2011-01-01T00:00:00Z', '2011-01-01T00:00:10Z', '2011-01-01T00:00:20Z'
    rows = [{'case_id': 'c1', 'event_id': 'e1', 'activity': 'A', 'timestamp': t0, 'lifecycle': 'complete'},
            {'case_id': 'c1', 'event_id': 'e2', 'activity': 'B', 'timestamp': t1, 'lifecycle': 'complete'},
            {'case_id': 'c1', 'event_id': 'e3', 'activity': 'C', 'timestamp': t1, 'lifecycle': 'complete'},
            {'case_id': 'c1', 'event_id': 'e4', 'activity': 'D', 'timestamp': t2, 'lifecycle': 'complete'}]
    swapped = [rows[0], rows[2], rows[1], rows[3]]
    g.equal(raw_dfg(rows), {'A->B': 1, 'B->C': 1, 'C->D': 1})
    g.equal(raw_dfg(swapped), {'A->C': 1, 'C->B': 1, 'B->D': 1})  # same events, different edges
    g.rejects(g.Blocked, lambda: checked_dfg(rows))
    first, second = checked_dfg(rows, tiebreak='event_id'), checked_dfg(swapped, tiebreak='event_id')
    g.equal(first, second); g.equal(first['edges'], {'A->B': 1, 'B->C': 1, 'C->D': 1}); g.equal(first['basis'], 'observation')
    with_start = rows + [{'case_id': 'c1', 'event_id': 'e0', 'activity': 'A', 'timestamp': '2010-12-31T23:59:50Z', 'lifecycle': 'start'}]
    g.equal('A->A' in raw_dfg(with_start), True)  # unfiltered lifecycle rows create a self-loop
    filtered = checked_dfg(with_start, tiebreak='event_id')
    g.equal(filtered['events_excluded_by_lifecycle'], 1); g.equal('A->A' in filtered['edges'], False)
    g.rejects(g.Blocked, lambda: checked_dfg(rows + [rows[0]], tiebreak='event_id'))
    return {'row_order_dependence_reproduced': True, 'tie_without_contract': 'blocked', 'lifecycle_self_loop_reproduced': True,
            'two_engine_agreement_on_declared_order': True, 'original_chain_rerun': False}


@case(24)
def chart_payload_and_asset_contract():
    rows = [{'activity': 'Confirmation of receipt', 'mean_hours': None, 'events': 1434},
            {'activity': 'T02 Check confirmation of receipt', 'mean_hours': 27.94, 'events': 1368},
            {'activity': 'T03 Adjust confirmation of receipt', 'mean_hours': 129.66, 'events': 55}]
    series_fig = go.Figure(go.Bar(x=[r['activity'] for r in rows], y=pd.Series([r['mean_hours'] for r in rows])))
    encoded = json.loads(pio.to_json(series_fig))['data'][0]['y']
    g.equal(sorted(encoded), ['bdata', 'dtype'])  # the pandas path stores a typed binary array
    decoded = np.frombuffer(base64.b64decode(encoded['bdata']), dtype=encoded['dtype'])
    g.equal(bool(np.isnan(decoded[0])), True)  # the missing value became NaN inside the payload
    cdn_page = pio.to_html(series_fig, include_plotlyjs='cdn', full_html=True)
    g.equal([u.startswith('https://cdn.plot.ly/') for u in external_assets(cdn_page)], [True])
    g.rejects(g.Blocked, lambda: chart_page(rows, include_plotlyjs='cdn'))
    result = chart_page(rows, include_plotlyjs=True)
    g.equal(result['missing_values_kept'], 1); g.require(result['bytes'] > 3_000_000, 'plotly.js not embedded')
    g.rejects(ValidationError, lambda: chart_page(rows + [{'activity': 'x', 'mean_hours': '5', 'events': 1}], include_plotlyjs=True))
    return {'cdn_dependency_reproduced': True, 'none_to_nan_in_binary_payload_reproduced': True, 'list_payload_keeps_null': True,
            'payload_equals_source_rows': True, 'counts_carried_per_bar': True, 'original_chain_rerun': False}


@case(25)
def required_action_rows_survive_the_document():
    rows = [{'fact_id': 'usc5-552-doj:u0089:s02#e53', 'actor': 'agency', 'modality': 'obligatory', 'action': 'make',
             'condition': 'in accordance with published rules', 'quote': 'agency, in accordance with published rules, shall make',
             'source_sentence': 'Each agency, in accordance with published rules, shall make available for public inspection.'},
            {'fact_id': 'nodejs-tsc-charter:u0001:s03#e3', 'actor': 'TSC', 'modality': 'obligatory', 'action': 'have',
             'condition': 'none stated', 'quote': 'must have', 'source_sentence': 'The TSC must have a Chair.'}]
    data = action_document(rows)
    g.equal(bool(DeepDiff(rows, read_action_document(data), zip_ordered_iterables=True)), False)
    g.rejects(ValidationError, lambda: action_document([dict(rows[1], actor='')]))
    g.rejects(ValidationError, lambda: action_document([dict(rows[0], quote='must make')]))
    g.rejects(UndefinedError, lambda: RENDERING.render(**{k: v for k, v in rows[0].items() if k != 'condition'}))
    d = docx.Document(io.BytesIO(data)); d.tables[0].rows[1].cells[1].text = ''
    buf = io.BytesIO(); d.save(buf); g.rejects(ValidationError, lambda: read_action_document(buf.getvalue()))
    d = docx.Document(io.BytesIO(data)); d.tables[0].rows[1].cells[4].text = 'none stated'
    buf = io.BytesIO(); d.save(buf); g.rejects(g.Blocked, lambda: read_action_document(buf.getvalue()))
    return {'typed_rows_roundtrip_native_docx': True, 'empty_actor': 'blocked', 'quote_outside_sentence': 'blocked',
            'missing_condition_field': 'blocked', 'edited_cell_detected': True, 'original_chain_rerun': False}


@case(26)
def lexical_candidates_are_not_authorizations():
    g.equal([hex(ord(c)) for c in RECORDED_P26_TOKEN], ['0x49', '0x74', '0xc3', '0xa2', '0xc2', '0x80', '0xc2', '0x99', '0x73'])
    g.equal(ftfy.badness.is_bad(RECORDED_P26_TOKEN), True)
    fixed = ftfy.fix_and_explain(RECORDED_P26_TOKEN)
    g.equal(fixed.text, "It's"); g.equal([s for s in fixed.explanation if s[0] == 'decode'], [('decode', 'utf-8'), ('decode', 'utf-8')])
    g.rejects(g.Blocked, lambda: checked_text(RECORDED_P26_TOKEN + ' not causing any maintenance burdens'))
    g.equal(checked_text('It’s not causing any maintenance burdens'), 'It’s not causing any maintenance burdens')
    tags = tag_candidates('the tsc discussed the meeting; remedy conflict', ACTIONS)
    g.equal(tags['overlap_winners'], sorted(n for n in ACTIONS if n != 'remedy'))  # equal-weight overlap: a seven-way tie on the, tsc, meeting
    g.equal(tags['ranked'][0]['action'], 'remedy'); g.equal(tags['state'], 'candidate')
    g.equal(tag_candidates('the tsc meeting', ACTIONS)['state'], 'tie')
    try:
        from nltk.corpus import wordnet as wn
        sense = wn.synset('approve.v.01')
        forms = [l.name().replace('_', ' ') for l in sense.lemmas()]
        antonyms = sorted({a.name() for l in sense.lemmas() for a in l.antonyms()})
        lexicon = {'APPROVE': forms}
        g.equal(g.resolve_alias('sanction the plan', lexicon, scope='review', allowed_scope='review', polarity='positive', expected_polarity='positive'), 'APPROVE')
        g.equal(g.alias_candidates('disapprove the plan', lexicon), [])
        g.equal(antonyms, ['disapprove'])
        located = nltk.data.find('corpora/wordnet.zip')
        corpus = getattr(located, 'path', None) or getattr(getattr(located, 'zipfile', None), 'filename', None)
        wordnet = {'version': wn.get_version(), 'sense': 'approve.v.01', 'forms': forms, 'antonyms': antonyms,
                   'corpus_sha256': hashlib.sha256(Path(corpus).read_bytes()).hexdigest() if corpus else 'unknown'}
    except LookupError:
        wordnet = {'state': 'not installed; no lexical expansion executed'}
    return {'recorded_mojibake_reproduced_and_blocked': True, 'ftfy_restores_two_decode_rounds': True,
            'overlap_tie_among_frequent_token_actions': 7, 'idf_weighted_winner': 'remedy', 'candidates_never_approved': True, 'wordnet': wordnet, 'original_chain_rerun': False}


@case(27)
def conformance_groups_need_full_coverage():
    record_rows = [{'case_id': 'a', 'department': 'General'}, {'case_id': 'b', 'department': 'General'}, {'case_id': 'c', 'department': 'General'}]
    conformance = [{'case_id': 'a', 'fitness': 1.0, 'state': 'fit'}, {'case_id': 'b', 'fitness': 0.5, 'state': 'deviates'}]
    partition = [{'case_id': 'a', 'split': 'train'}, {'case_id': 'b', 'split': 'holdout'}, {'case_id': 'c', 'split': 'holdout'}]
    inner = pd.DataFrame(record_rows).merge(pd.DataFrame(conformance), on='case_id')
    g.equal(len(inner), 2)  # the default inner join drops the unassessed case silently
    g.rejects(g.Blocked, lambda: conformance_by_group(record_rows, conformance, partition))
    complete = conformance + [{'case_id': 'c', 'fitness': None, 'state': 'not_run'}]
    out = conformance_by_group(record_rows, complete, partition)
    g.equal(out['groups'], [{'department': 'General', 'split': 'holdout', 'cases': 2, 'assessed': 1, 'fit': 0, 'not_run': 1, 'mean_fitness': 0.5},
                            {'department': 'General', 'split': 'train', 'cases': 1, 'assessed': 1, 'fit': 1, 'not_run': 0, 'mean_fitness': 1.0}])
    g.equal(out['regulatory_compliance_established'], False)
    g.rejects(g.Blocked, lambda: conformance_by_group(record_rows, complete, partition[:2]))
    g.rejects(g.Blocked, lambda: conformance_by_group(record_rows, complete[:2] + [{'case_id': 'c', 'fitness': 1.0, 'state': 'timeout_is_fit'}], partition))
    return {'inner_join_loss_reproduced': True, 'coverage_mismatch': 'blocked', 'training_cases_separated': True,
            'unassessed_excluded_from_mean': True, 'original_chain_rerun': False}


@case(28)
def meeting_join_keeps_tag_states():
    hostile = 'm"). :- a. %'
    decisions = [{'meeting': 'tsc-2023-11-08', 'majority_reachable': False}, {'meeting': 'tsc-2023-12-06', 'majority_reachable': True},
                 {'meeting': 'tsc-2024-01-10', 'majority_reachable': True}, {'meeting': hostile, 'majority_reachable': True}]
    tags = [{'meeting': 'tsc-2023-12-06', 'action': 'open', 'state': 'accepted'}, {'meeting': 'tsc-2023-12-06', 'action': 'have', 'state': 'accepted'},
            {'meeting': 'tsc-2024-01-10', 'action': 'approve', 'state': 'tie'}, {'meeting': 'tsc-2024-01-10', 'action': 'do', 'state': 'tie'},
            {'meeting': 'tsc-2099-01-01', 'action': 'take', 'state': 'accepted'}, {'meeting': hostile, 'action': 'x', 'state': 'accepted'}]
    out = meeting_join(decisions, tags)
    g.equal(out['discussed'], {'tsc-2023-11-08': 0, 'tsc-2023-12-06': 2, 'tsc-2024-01-10': 0, hostile: 1})
    g.equal(out['tied'], {'tsc-2023-11-08': 0, 'tsc-2023-12-06': 0, 'tsc-2024-01-10': 2, hostile: 0})
    g.equal(out['untagged'], ['tsc-2023-11-08']); g.equal(out['unmatched'], ['tsc-2099-01-01'])
    g.rejects(g.Blocked, lambda: meeting_join(decisions + [{'meeting': 'tsc-2023-12-06', 'majority_reachable': False}], tags))
    g.rejects(g.Blocked, lambda: meeting_join(decisions, tags + [{'meeting': 'tsc-2023-12-06', 'action': 'y', 'state': 'maybe'}]))
    return {'tied_and_untagged_distinguished_from_zero': True, 'unmatched_meeting_reported': True, 'hostile_id_stays_one_literal': True,
            'conflicting_attendance_rows': 'blocked', 'attendance_is_a_scenario_not_a_vote': True, 'original_chain_rerun': False}


@case(29)
def execution_witness_versus_structural_precedence():
    merge_trace = run_witness(BPMN_MERGE, 'proc')
    g.equal(merge_trace, ['a', 'b', 'c', 'c'])  # uncontrolled merge: c runs once per arriving token, and the workflow still completes
    g.rejects(g.Blocked, lambda: execution_checked(BPMN_MERGE, 'proc', [('a', 'c')]))
    out = execution_checked(BPMN_JOIN, 'proc', [('a', 'c'), ('b', 'c')])
    g.equal(out['witness_trace'], ['a', 'b', 'c'])
    g.require(out['witness_trace'].index('a') < out['witness_trace'].index('b'), 'name order witness')
    g.rejects(g.Blocked, lambda: execution_checked(BPMN_JOIN, 'proc', [('a', 'b')]))  # witness order is not a forced edge
    g.rejects(g.Blocked, lambda: execution_checked(BPMN_JOIN, 'proc', [('a', 'zz')]))
    graph = sequence_flow_graph(BPMN_JOIN)
    g.equal(graph.nodes['join']['kind'], 'parallelGateway'); g.equal(nx.is_directed_acyclic_graph(graph), True)
    return {'double_execution_under_uncontrolled_merge_reproduced': True, 'completion_flag_insufficient': True,
            'forced_edges_checked_by_has_path': True, 'scheduling_order_not_evidence': True, 'original_chain_rerun': False}


@case(30)
def workbook_readback_typed_diff():
    rows = [{'id': 'r1', 'count': 1, 'ratio': 1.5, 'code': '001', 'flag': True, 'missing': None, 'formula_like': '=1+1', 'numeric_text': '1.0'}]
    typed = workbook_roundtrip(rows, profile='typed')
    changes = typed['diff']['type_changes']
    g.equal(changes["root[0]['count']"]['new_type'], 'float'); g.equal(changes["root[0]['missing']"], {'old_type': 'NoneType', 'new_type': 'str', 'old_value': None, 'new_value': ''})
    g.equal(typed['diff']['values_changed']["root[0]['formula_like']"], {'new_value': '', 'old_value': '=1+1'})
    g.equal(set(typed['diff']) & {'code', 'flag'}, set())
    g.rejects(g.Blocked, lambda: checked_roundtrip(rows, profile='typed'))
    g.equal(checked_roundtrip(rows, profile='json_text')['diff'], {})
    g.equal(csv_compare({'r1': {'id': 'r1', 'v': 1}}, {'r1': {'id': 'r1', 'v': True}})['changed'], [])
    g.equal(list(DeepDiff({'v': 1}, {'v': True})), ['type_changes'])
    g.rejects(g.Blocked, lambda: workbook_roundtrip(rows, profile='native'))
    return {'int_to_float_reproduced': True, 'none_to_empty_string_reproduced': True, 'formula_like_text_lost_reproduced': True,
            'csv_diff_bool_int_conflation_reproduced': True, 'declared_json_text_profile_roundtrips': True, 'original_chain_rerun': False}


@case(0)
def library_primitive_substitutions_for_earlier_glue():
    expected = {'start': 'start', 'choose': 'exclusive', 'approve': 'task', 'end': 'end'}
    rows = [{'id': k, 'kind': v} for k, v in expected.items()]
    fixtures = [rows, rows[:-1], rows + [rows[0]], [dict(r, kind='parallel') if r['id'] == 'choose' else r for r in rows]]
    g.equal([verdict(g.validate_process_map, expected, f) for f in fixtures], [verdict(process_map_diff, expected, f) for f in fixtures])
    cells = [None, '', False, 0, '001', '=1+1']
    variants = [list(cells)] + [list(cells[:i]) + [v] + list(cells[i + 1:]) for i, v in [(0, 0), (1, None), (2, 0), (4, 1), (5, 2)]]
    earlier = [verdict(g2.typed_readback, cells, v) for v in variants]
    positional = ['accepted' if not DeepDiff(cells, v, zip_ordered_iterables=True) else 'rejected' for v in variants]
    default_pairing = ['accepted' if not DeepDiff(cells, v) else 'rejected' for v in variants]
    g.equal(earlier, positional)
    g.equal(sum(a == b for a, b in zip(earlier, default_pairing)), 5)  # default list pairing aligns False with 0 and hides the type change
    roster = {'valid_from': date(2020, 1, 1), 'valid_to': date(2021, 1, 1), 'coverage': 'complete', 'voters': ['a', 'b', 'c', 'd']}
    probes = [(date(2020, 3, 1), [roster]), (date(2019, 1, 1), [roster]), (date(2021, 1, 1), [roster]), (date(2020, 3, 1), [roster, roster])]
    g.equal([verdict(g2.majority_at, m, r, {'a'}) for m, r in probes], ['accepted' if len(rosters_at(m, r)) == 1 else 'rejected' for m, r in probes])
    receipt = {'case_id': 'c', 'action_id': 'review', 'receipt_id': 'r', 'timestamp': '2026-01-01T00:00:00Z'}
    evidence = [('observed_completion', receipt), ('proposal', receipt), ('observed_completion', {'case_id': 'c'}), ('observed_completion', dict(receipt, case_id=''))]
    g.equal([verdict(g.observed_action, k, e) for k, e in evidence], [verdict(OccurrenceReceipt.model_validate, {'kind': k, **e}) for k, e in evidence])
    return {'validate_process_map': 'deepdiff.DeepDiff + pandas.Index.has_duplicates, 4 of 4 verdicts agree',
            'typed_readback': 'deepdiff.DeepDiff(zip_ordered_iterables=True), 6 of 6 verdicts agree; default pairing agrees on 5 of 6', 'majority_at interval selection': 'portion.closedopen containment, 4 of 4 verdicts agree',
            'observed_action': 'pydantic strict model with Literal kind, 4 of 4 verdicts agree'}


def run() -> dict:
    g.ASSERTIONS = 0; rows = []
    for i, fn in C:
        before = g.ASSERTIONS
        try: r = fn(); state = 'passed'; error = None
        except Exception: r = {}; state = 'failed'; error = traceback.format_exc()
        rows.append({'proof': i, 'test': fn.__name__, 'status': state, 'assertions': g.ASSERTIONS - before, 'result': r, 'error': error})
    packages = ['pydantic', 'pandas', 'numpy', 'duckdb', 'portion', 'deepdiff', 'ftfy', 'scikit-learn', 'pm4py', 'plotly', 'lxml',
                'python-docx', 'openpyxl', 'python-calamine', 'csv-diff', 'SpiffWorkflow', 'networkx', 'clingo', 'nltk', 'Jinja2', 'spacy']
    here = Path(__file__).resolve().parent
    return {'scope': 'Bounded authored handoff fixtures and reproduced library defaults; the original P21-P30 chains and their artifacts were not rerun.',
            'test_cases': len(rows), 'passed': sum(r['status'] == 'passed' for r in rows), 'failed': sum(r['status'] == 'failed' for r in rows),
            'assertions': g.ASSERTIONS, 'full_chains_executed': 0, 'network_calls': 0, 'business_actions': 0,
            'source_sha256': hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
            'dependency_sha256': {name: hashlib.sha256((here / name).read_bytes()).hexdigest() for name in ['handoff_guards_v1.py', 'handoff_guards_v2.py']},
            'library_versions': {n: importlib.metadata.version(n) for n in packages}, 'results': rows}


if __name__ == '__main__':
    p = argparse.ArgumentParser(); p.add_argument('--report', type=Path, required=True); args = p.parse_args()
    result = run(); args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(json.dumps(result, indent=2, sort_keys=True, allow_nan=False) + '\n')
    print(json.dumps({k: result[k] for k in ['test_cases', 'passed', 'failed', 'assertions', 'full_chains_executed']}))
    if result['failed']:
        print('\n'.join(r['error'] for r in result['results'] if r['error'])); raise SystemExit(1)
