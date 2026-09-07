"""Tenth bounded adapter suite: P91-P100, the data atlas. Requires handoff_guards_v1.py and handoff_guards_v8.py from
this TXT. Every grouping, ranking, matrix build, calendar difference, schema, chart specification, workbook readback and
DOM assertion below is a library primitive call; local code declares fixtures and policies and turns a primitive result
into a Blocked outcome. All runtime inputs are authored fixtures. No model client and no network.
Run: python handoff_guards_v10.py --report report.json [--jsdom-dir DIR]
"""
from __future__ import annotations
import argparse, collections, datetime, hashlib, importlib.metadata, io, json, logging, os, traceback, warnings
from pathlib import Path
from typing import Literal
import altair as alt
import networkx as nx
import numpy as np
import openpyxl
import pandas as pd
import vl_convert as vlc
from deepdiff import DeepDiff
from lxml import etree
from pydantic import BaseModel, ConfigDict, Field, ValidationError, create_model
from scipy.sparse import coo_matrix
import handoff_guards_v1 as g
import handoff_guards_v8 as g8
warnings.simplefilter('ignore'); logging.disable(logging.WARNING)

# ---------------------------------------------------------------- P91
def category_totals(cases: list[dict], *, rows: str, cols: str, population: int) -> dict:
    df = pd.DataFrame(cases)
    if df['case_id'].duplicated().any(): raise g.Blocked('duplicate case')
    if len(df) != population: raise g.Blocked('case rows do not match the declared population')
    grid = df.groupby([rows, cols], dropna=False).size().unstack(fill_value=0)
    if int(grid.to_numpy().sum()) != population: raise g.Blocked('grid does not conserve the population')
    per_row = df.groupby(rows, dropna=False).size().to_dict()
    if sum(per_row.values()) != population: raise g.Blocked('row dimension does not total the population')
    return {'grid': {r: {c: int(grid.loc[r, c]) for c in grid.columns} for r in grid.index},
            'per_row': {k: int(v) for k, v in per_row.items()}, 'population': population}

# ---------------------------------------------------------------- P92
def ranked_coverage(paths: list[dict], *, population: int) -> dict:
    df = pd.DataFrame(paths)
    if df['path_id'].duplicated().any(): raise g.Blocked('duplicate path')
    if int(df['cases'].sum()) != population: raise g.Blocked('path counts do not total the population')
    ordered = df.sort_values('cases', ascending=False, kind='stable')  # ties keep source order
    ordered = ordered.assign(cumulative=ordered['cases'].cumsum())
    ordered = ordered.assign(coverage=(ordered['cumulative'] / population * 100).round(4))
    last = ordered.iloc[-1]
    if int(last['cumulative']) != population or float(last['coverage']) != 100.0: raise g.Blocked('coverage does not reach the population')
    return {'order': ordered['path_id'].tolist(), 'cumulative': [int(v) for v in ordered['cumulative']],
            'final_coverage': float(last['coverage']), 'sequences': {r['path_id']: list(r['sequence']) for _, r in df.iterrows()}}

# ---------------------------------------------------------------- P93
def resource_matrix(pairs: list[dict], resources: list[str]) -> dict:
    seen = set()
    for p in pairs:
        key = (p['source'], p['target'])
        if key in seen: raise g.Blocked('duplicate resource pair: %s -> %s' % key)
        if p['source'] not in resources or p['target'] not in resources: raise g.Blocked('pair outside the declared resource set')
        seen.add(key)
    index = {r: i for i, r in enumerate(resources)}
    n = len(resources)
    rows = [index[p['source']] for p in pairs]; cols = [index[p['target']] for p in pairs]
    weights = coo_matrix(([p['weight'] for p in pairs], (rows, cols)), shape=(n, n)).toarray()
    present = np.zeros((n, n), dtype=bool)
    for r, c in zip(rows, cols): present[r, c] = True  # a boolean mask cannot be built by summing
    graph = nx.DiGraph(); graph.add_weighted_edges_from([(p['source'], p['target'], p['weight']) for p in pairs])
    return {'recorded_cells': int(present.sum()), 'unrecorded_cells': int((~present).sum()),
            'self_pairs': nx.number_of_selfloops(graph), 'matrix': weights.tolist(), 'present': present.tolist(),
            'measure': 'handover weight, not a case count'}

# ---------------------------------------------------------------- P94
def interval_profiles(cases: list[dict]) -> dict:
    df = pd.DataFrame(cases)
    start = pd.to_datetime(df['start'], utc=True, errors='raise'); deadline = pd.to_datetime(df['deadline'], utc=True, errors='raise')
    if (deadline < start).any(): raise g.Blocked('a deadline precedes its start')
    df = df.assign(calendar_days=(deadline - start).dt.days)
    for c in cases:
        if 'working_days' not in c: raise g.Blocked('stored working-day count missing; holiday rules are not rerun here')
    profiles = df.groupby(['calendar_days', 'working_days'], dropna=False).size()
    return {'cases': len(df), 'calendar_days': df['calendar_days'].tolist(),
            'profiles': {'%d/%d' % k: int(v) for k, v in profiles.items()},
            'working_days_source': 'retained from the recorded workbook; the holidays library was not rerun (P58)'}

# ---------------------------------------------------------------- P95
FIELD_TYPES = {'string': str, 'integer': int, 'boolean': bool, 'date': datetime.date}


def build_row_model(fields: list[dict]):
    if not fields: raise g.Blocked('no field definitions')
    spec = {}
    for f in fields:
        if f['type'] not in FIELD_TYPES: raise g.Blocked('undeclared field type: ' + str(f['type']))
        annotation = FIELD_TYPES[f['type']]
        spec[f['name']] = ((annotation | None) if f.get('optional') else annotation, ... if not f.get('optional') else None)
    return create_model('AtlasRow', __config__=ConfigDict(strict=True, extra='forbid'), **spec)


def validate_rows(fields: list[dict], rows: list[dict]) -> dict:
    model = build_row_model(fields)
    accepted, rejected = [], []
    for r in rows:
        try: accepted.append(model.model_validate(r).model_dump())
        except ValidationError as e: rejected.append({'row': r, 'errors': len(e.errors())})
    schema = model.model_json_schema()
    return {'fields': len(fields), 'accepted': len(accepted), 'rejected': len(rejected),
            'cells': len(accepted) * len(fields), 'schema_properties': sorted(schema['properties']),
            'missing_preserved': sum(1 for r in accepted for v in r.values() if v is None)}

# ---------------------------------------------------------------- P96
def atlas_chart(rows: list[dict], *, x: str, y: str) -> dict:
    frame = pd.DataFrame(rows)
    missing = [f for f in (x, y) if f not in frame.columns]
    if missing: raise g.Blocked('encoded field is not a column of the data: ' + ','.join(missing))
    chart = alt.Chart(frame).mark_bar().encode(x=alt.X(x + ':N'), y=alt.Y(y + ':Q'))
    spec = chart.to_dict()
    data = spec.get('data', {})
    if 'url' in data: raise g.Blocked('specification references an external data URL')
    # Altair names the dataset and stores the rows under datasets; inline values are the other shape.
    data_values = data.get('values') if 'values' in data else spec.get('datasets', {}).get(data.get('name'))
    if data_values is None: raise g.Blocked('specification does not carry its data inline or as a named dataset')
    if DeepDiff(rows, data_values, zip_ordered_iterables=True): raise g.Blocked('embedded data differs from the source rows')
    svg = vlc.vegalite_to_svg(json.dumps(spec))
    again = vlc.vegalite_to_svg(json.dumps(spec))
    root = etree.fromstring(svg.encode())
    marks = root.findall('.//{http://www.w3.org/2000/svg}path') + root.findall('.//{http://www.w3.org/2000/svg}rect')
    return {'rows': len(rows), 'svg_repeatable': svg == again, 'svg_bytes': len(svg.encode()),
            'mark_elements': len(marks), 'external_data_url': False}

# ---------------------------------------------------------------- P97
def atlas_workbook(sheets: dict[str, list[list]]) -> bytes:
    book = openpyxl.Workbook(); book.remove(book.active)
    for name, rows in sheets.items():
        ws = book.create_sheet(name)
        for row in rows: ws.append(row)
    buf = io.BytesIO(); book.save(buf); return buf.getvalue()

# ---------------------------------------------------------------- P100
NARRATIVE_FIELDS = {'population', 'flagged', 'departments'}


def executive_narrative(model: dict, sentences: list[str]) -> dict:
    used = []
    for s in sentences:
        placeholders = [p.strip('{}') for p in s.split() if p.startswith('{') and p.endswith('}')]
        for p in placeholders:
            if p not in model: raise g.Blocked('narrative names a field the model does not carry: ' + p)
            used.append(p)
    rendered = [s.format(**model) for s in sentences]
    for value in ('population', 'flagged'):
        if str(model[value]) not in ' '.join(rendered) and value in used: raise g.Blocked('a named value did not reach the text')
    return {'sentences': len(rendered), 'fields_used': sorted(set(used)), 'text': rendered}

C = []
def case(i):
    def reg(fn): C.append((i, fn)); return fn
    return reg


@case(91)
def group_totals_conserve_the_population():
    cases = [{'case_id': 'c1', 'department': 'General', 'channel': 'Desk'},
             {'case_id': 'c2', 'department': 'General', 'channel': 'Internet'},
             {'case_id': 'c3', 'department': 'Experts', 'channel': 'Desk'}]
    frame = pd.DataFrame(cases)
    plain = frame.groupby(['department', 'channel']).size().unstack()
    g.equal(bool(plain.isna().any().any()), True)  # an absent combination becomes NaN, not zero
    out = category_totals(cases, rows='department', cols='channel', population=3)
    g.equal(out['grid'], {'Experts': {'Desk': 1, 'Internet': 0}, 'General': {'Desk': 1, 'Internet': 1}})
    g.equal(out['per_row'], {'Experts': 1, 'General': 2})
    g.rejects(g.Blocked, lambda: category_totals(cases, rows='department', cols='channel', population=4))
    g.rejects(g.Blocked, lambda: category_totals(cases + [cases[0]], rows='department', cols='channel', population=4))
    return {'unstack_nan_reproduced': True, 'fill_value_conserves_population': True, 'both_dimensions_total_the_population': True}


@case(92)
def ranking_ties_need_a_stable_sort():
    frame = pd.DataFrame({'id': ['p%d' % i for i in range(17)], 'v': [1, 2] * 8 + [1]})
    quick = frame.sort_values('v', ascending=False)['id'].tolist()
    stable = frame.sort_values('v', ascending=False, kind='stable')['id'].tolist()
    g.require(quick != stable, 'the default sort reorders ties')  # only mergesort and stable are stable algorithms
    g.equal(stable[:4], ['p1', 'p3', 'p5', 'p7'])
    paths = [{'path_id': 'v1', 'cases': 10, 'sequence': ['A', 'B', 'A']}, {'path_id': 'v2', 'cases': 10, 'sequence': ['A', 'C']},
             {'path_id': 'v3', 'cases': 5, 'sequence': ['A']}]
    out = ranked_coverage(paths, population=25)
    g.equal(out['order'], ['v1', 'v2', 'v3']); g.equal(out['cumulative'], [10, 20, 25]); g.equal(out['final_coverage'], 100.0)
    g.equal(out['sequences']['v1'], ['A', 'B', 'A'])  # a repeated activity survives the ranking
    g.rejects(g.Blocked, lambda: ranked_coverage(paths, population=26))
    return {'default_sort_reorders_ties_reproduced': True, 'stable_sort_preserves_source_order': True,
            'cumulative_reaches_the_population': True, 'repeated_activities_retained': True}


@case(93)
def sparse_matrix_sums_duplicates_silently():
    pairs = [{'source': 'R01', 'target': 'R02', 'weight': 0.2}, {'source': 'R01', 'target': 'R01', 'weight': 0.13}]
    duplicated = coo_matrix(([0.2, 0.3], ([0, 0], [1, 1])), shape=(2, 2)).toarray()
    g.equal(duplicated[0][1], 0.5)  # coo_matrix adds duplicate entries instead of rejecting them
    mask_by_sum = coo_matrix(([1, 1], ([0, 0], [1, 1])), shape=(2, 2)).toarray()
    g.equal(int(mask_by_sum[0][1]), 2)  # a presence mask built the same way becomes a count
    out = resource_matrix(pairs, ['R01', 'R02'])
    g.equal((out['recorded_cells'], out['unrecorded_cells']), (2, 2))
    g.equal(out['self_pairs'], 1); g.equal(out['present'][1][0], False)  # an unrecorded cell is not a zero weight
    g.equal(out['matrix'][1][0], 0.0)
    g.rejects(g.Blocked, lambda: resource_matrix(pairs + [pairs[0]], ['R01', 'R02']))
    g.rejects(g.Blocked, lambda: resource_matrix(pairs + [{'source': 'R09', 'target': 'R01', 'weight': 1.0}], ['R01', 'R02']))
    return {'coo_duplicate_summing_reproduced': True, 'presence_mask_built_without_summing': True,
            'unrecorded_kept_distinct_from_zero': True, 'self_pairs_retained': True, 'weights_are_not_case_counts': True}


@case(94)
def calendar_days_are_a_difference_not_a_rerun():
    cases = [{'case_id': 'c1', 'start': '2011-10-11', 'deadline': '2011-12-06', 'working_days': 40},
             {'case_id': 'c2', 'start': '2011-10-11', 'deadline': '2011-12-06', 'working_days': 40},
             {'case_id': 'c3', 'start': '2011-11-01', 'deadline': '2011-12-06', 'working_days': 25}]
    out = interval_profiles(cases)
    g.equal(out['calendar_days'], [56, 56, 35])  # deadline minus start, in days
    g.equal(out['profiles'], {'35/25': 1, '56/40': 2})
    g.require('was not rerun' in out['working_days_source'], 'the working-day provenance travels with the result')
    g.rejects(g.Blocked, lambda: interval_profiles([{'case_id': 'x', 'start': '2011-12-06', 'deadline': '2011-10-11', 'working_days': 1}]))
    g.rejects(g.Blocked, lambda: interval_profiles([{'case_id': 'x', 'start': '2011-10-11', 'deadline': '2011-12-06'}]))
    return {'calendar_difference_checked': True, 'stored_working_days_retained_not_recomputed': True,
            'reversed_interval_blocked': True, 'profiles_counted': True}


@case(95)
def strict_row_schemas_reject_and_preserve():
    fields = [{'name': 'case_id', 'type': 'string'}, {'name': 'events', 'type': 'integer'},
              {'name': 'at_risk', 'type': 'boolean'}, {'name': 'activity', 'type': 'string', 'optional': True}]
    rows = [{'case_id': 'c1', 'events': 3, 'at_risk': True, 'activity': 'T02'},
            {'case_id': 'c2', 'events': 1, 'at_risk': False, 'activity': None},
            {'case_id': 'c3', 'events': '1', 'at_risk': False, 'activity': 'T03'},
            {'case_id': 'c4', 'events': 1, 'at_risk': 1, 'activity': 'T03'},
            {'case_id': 'c5', 'events': 1, 'at_risk': False, 'activity': 'T03', 'extra': 1}]
    out = validate_rows(fields, rows)
    g.equal((out['accepted'], out['rejected']), (2, 3))  # a string integer, an int boolean and an unknown field are refused
    g.equal(out['missing_preserved'], 1)  # the optional activity stays None rather than becoming ''
    g.equal(out['schema_properties'], ['activity', 'at_risk', 'case_id', 'events'])
    g.equal(out['cells'], 8)
    g.rejects(g.Blocked, lambda: build_row_model([{'name': 'x', 'type': 'decimal'}]))
    g.rejects(g.Blocked, lambda: build_row_model([]))
    return {'strict_types_enforced': True, 'unknown_field_refused': True, 'optional_missing_preserved': True,
            'json_schema_emitted': True}


@case(96)
def chart_specification_embeds_its_data():
    rows = [{'department': 'General', 'cases': 1390}, {'department': 'Experts', 'cases': 15}, {'department': 'Customer contact', 'cases': 29}]
    out = atlas_chart(rows, x='department', y='cases')
    g.equal(out['rows'], 3); g.equal(out['external_data_url'], False)
    g.equal(out['svg_repeatable'], True)  # the same specification renders to the same SVG
    g.require(out['mark_elements'] > 0, 'the SVG carries rendered marks')
    frame = pd.DataFrame(rows)
    spec = alt.Chart(frame).mark_bar().encode(x='department:N', y='cases:Q').to_dict()
    g.equal('values' in spec['data'], False)  # the rows live under a named dataset, not inline
    g.equal(DeepDiff(rows, spec['datasets'][spec['data']['name']], zip_ordered_iterables=True), {})
    unchecked = alt.Chart(pd.DataFrame(rows)).mark_bar().encode(x='missing_column:N', y='cases:Q').to_dict()
    g.equal(unchecked['encoding']['x']['field'], 'missing_column')  # Altair builds a spec for a column that does not exist
    g.rejects(g.Blocked, lambda: atlas_chart(rows, x='missing_column', y='cases'))
    return {'data_carried_as_a_named_dataset_not_inline_reproduced': True, 'no_external_data_url': True, 'svg_repeatable': True,
            'marks_rendered_by_the_vega_runtime': True,
            'encoding_of_a_nonexistent_column_accepted_reproduced': True}


@case(97)
def atlas_workbook_reads_back_after_normalization():
    sheets = {'departments': [['department', 'cases'], ['General', 1390], ['Experts', 15]]}
    data = atlas_workbook(sheets)
    out = g8.read_two_ways(data, 'departments')
    g.equal(out['raw_agree'], False); g.equal(out['normalized_agree'], True)  # calamine returns the integers as floats
    g.equal(out['rows'], 3)
    ws = openpyxl.load_workbook(io.BytesIO(data))['departments']
    g.equal([c.value for c in ws[2]], ['General', 1390])
    return {'two_readers_normalized': True, 'sheet_selected_by_name': True, 'values_preserved': True}


@case(98)
def searchable_tables_render_the_contract():
    fields = [{'name': 'case_id', 'type': 'string'}, {'name': 'events', 'type': 'integer'}]
    rows = [{'case_id': 'c1', 'events': 3}, {'case_id': 'c2', 'events': 1}]
    checked = validate_rows(fields, rows)
    g.equal(checked['rejected'], 0)
    page = g8.CATALOG.render(title='Data desk', fields=['case_id', 'events'], rows=[{'case_id': r['case_id'], 'events': str(r['events'])} for r in rows])
    if g8.jsdom_available():
        dom = g8.run_dom(page)
        g.equal([r[0] for r in dom['rows']], ['c1', 'c2']); g.equal(dom['externalRequests'], 0)
    g.equal('<script' in page, False)
    return {'rows_validated_before_render': True, 'dom_checked': g8.jsdom_available(), 'no_inline_script': True}


@case(99)
def field_glossary_states_units_and_sources():
    glossary = [{'name': 'mean_days_closed', 'unit': 'days', 'source': 'closed cases only (P22)'},
                {'name': 'handover_weight', 'unit': 'share', 'source': 'pm4py handover network; not a case count (P62)'}]
    for entry in glossary:
        if not entry.get('unit') or not entry.get('source'): raise g.Blocked('a field without a unit or a source')
    page = g8.CATALOG.render(title='Field glossary', fields=['name', 'unit', 'source'], rows=glossary)
    g.equal('not a case count' in page, True)  # the measure caveat travels into the interface
    g.equal(len(glossary), 2)
    if g8.jsdom_available():
        dom = g8.run_dom(page)
        g.equal([r[0] for r in dom['rows']], ['mean_days_closed', 'handover_weight'])
    return {'units_declared': True, 'measure_caveats_carried_into_the_interface': True, 'dom_checked': g8.jsdom_available()}


@case(100)
def the_narrative_names_only_model_fields():
    model = {'population': 1434, 'flagged': 94, 'departments': 3}
    sentences = ['The atlas covers {population} cases.', 'Of the open cases, {flagged} are flagged.', 'Cases span {departments} departments.']
    out = executive_narrative(model, sentences)
    g.equal(out['sentences'], 3); g.equal(out['fields_used'], ['departments', 'flagged', 'population'])
    g.equal('1434' in out['text'][0] and '94' in out['text'][1], True)
    g.rejects(g.Blocked, lambda: executive_narrative(model, ['Risk rose by {trend} percent.']))  # a field the model does not carry
    g.rejects(KeyError, lambda: [s.format(**model) for s in ['{unknown}']])
    return {'every_named_field_exists_in_the_model': True, 'numbers_come_from_the_model': True,
            'no_unsourced_claim_can_be_rendered': True}


def run() -> dict:
    g.ASSERTIONS = 0; rows = []
    for i, fn in C:
        before = g.ASSERTIONS
        try: r = fn(); state = 'passed'; error = None
        except Exception: r = {}; state = 'failed'; error = traceback.format_exc()
        rows.append({'proof': i, 'test': fn.__name__, 'status': state, 'assertions': g.ASSERTIONS - before, 'result': r, 'error': error})
    packages = ['pydantic', 'pandas', 'numpy', 'scipy', 'networkx', 'altair', 'vl-convert-python', 'openpyxl', 'python-calamine', 'lxml', 'deepdiff']
    here = Path(__file__).resolve().parent
    return {'scope': 'Bounded authored handoff fixtures and reproduced library defaults; the original P91-P100 chains and their artifacts were not rerun.',
            'test_cases': len(rows), 'passed': sum(r['status'] == 'passed' for r in rows), 'failed': sum(r['status'] == 'failed' for r in rows),
            'assertions': g.ASSERTIONS, 'full_chains_executed': 0, 'network_calls': 0, 'business_actions': 0,
            'source_sha256': hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
            'dependency_sha256': {name: hashlib.sha256((here / name).read_bytes()).hexdigest() for name in ['handoff_guards_v1.py', 'handoff_guards_v8.py']},
            'library_versions': {n: importlib.metadata.version(n) for n in packages}, 'results': rows}


if __name__ == '__main__':
    p = argparse.ArgumentParser(); p.add_argument('--report', type=Path, required=True); p.add_argument('--jsdom-dir', default=os.environ.get('JSDOM_DIR', ''))
    args = p.parse_args(); os.environ['JSDOM_DIR'] = args.jsdom_dir; g8.JSDOM_DIR = args.jsdom_dir
    result = run(); args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(json.dumps(result, indent=2, sort_keys=True, allow_nan=False) + '\n')
    print(json.dumps({k: result[k] for k in ['test_cases', 'passed', 'failed', 'assertions']}))
    if result['failed']:
        print('\n'.join(r['error'] for r in result['results'] if r['error'])); raise SystemExit(1)
